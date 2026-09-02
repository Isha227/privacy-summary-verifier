from __future__ import annotations

import json
import re
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pandas as pd


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data" / "v2_prompt_intervention_v1_2" / "source_unit_method_validation"
ADJ = BASE / "v3_researcher_comparison" / "adjudication"
WORKBOOK = ADJ / "source_unit_comparison_v3_ADJUDICATED.xlsx"
UNANIMOUS = ADJ / "unanimous_resolutions.csv"
OUTPUT_CSV = ADJ / "final_v3_source_unit_decisions.csv"
OUTPUT_JSON = ADJ / "v3_source_unit_comparison_results.json"

SPECS = {
    "Candidate disagreements": {
        "direction": "candidate_to_reference",
        "expected_rows": 10,
        "labels": {
            "Equivalent",
            "Partially overlaps",
            "Valid additional unit",
            "Invalid / not important",
        },
        "match_prefix": "B-",
        "requires_match": {"Equivalent", "Partially overlaps"},
        "forbids_match": {"Valid additional unit"},
    },
    "Reference disagreements": {
        "direction": "reference_to_candidate",
        "expected_rows": 11,
        "labels": {"Fully represented", "Partly represented", "Not represented"},
        "match_prefix": "A-",
        "requires_match": {"Fully represented", "Partly represented"},
        "forbids_match": {"Not represented"},
    },
}


def clean(value) -> str:
    return "" if value is None else str(value).strip()


def split_ids(value: str) -> list[str]:
    value = clean(value)
    if not value or value.lower() in {"none", "n/a", "na", "-"}:
        return []
    return [part.strip() for part in re.split(r"[;,]", value) if part.strip()]


def read_adjudicated() -> tuple[pd.DataFrame, list[str]]:
    wb = openpyxl.load_workbook(WORKBOOK, data_only=False)
    errors: list[str] = []
    rows: list[dict] = []

    for sheet_name, spec in SPECS.items():
        if sheet_name not in wb.sheetnames:
            errors.append(f"Missing sheet: {sheet_name}")
            continue
        ws = wb[sheet_name]
        headers = [clean(ws.cell(3, col).value) for col in range(1, 12)]
        expected_headers = [
            "Policy", "Unit ID", "Important information", "Source evidence",
            "Material qualifiers", "A1 decision", "A2 decision", "A3 decision",
            "Final matching IDs", "Final label", "Final reason",
        ]
        if headers != expected_headers:
            errors.append(f"Unexpected headers in {sheet_name}: {headers}")

        data_rows = list(range(4, ws.max_row + 1))
        if len(data_rows) != spec["expected_rows"]:
            errors.append(
                f"{sheet_name}: expected {spec['expected_rows']} rows, found {len(data_rows)}"
            )

        for row_no in data_rows:
            policy = clean(ws.cell(row_no, 1).value)
            unit_id = clean(ws.cell(row_no, 2).value)
            match_text = clean(ws.cell(row_no, 9).value)
            label = clean(ws.cell(row_no, 10).value)
            reason = clean(ws.cell(row_no, 11).value)
            matches = split_ids(match_text)

            if not policy or not unit_id:
                errors.append(f"{sheet_name} row {row_no}: blank policy or unit ID")
            if label not in spec["labels"]:
                errors.append(f"{unit_id}: invalid final label {label!r}")
            if not reason:
                errors.append(f"{unit_id}: blank final reason")
            if label in spec["requires_match"] and not matches:
                errors.append(f"{unit_id}: {label} requires at least one matching ID")
            if label in spec["forbids_match"] and matches:
                errors.append(f"{unit_id}: {label} should not have matching IDs")
            for match_id in matches:
                if not match_id.startswith(spec["match_prefix"]):
                    errors.append(f"{unit_id}: invalid matched-set ID {match_id}")
                if f"-{policy}-" not in match_id:
                    errors.append(f"{unit_id}: cross-policy matched ID {match_id}")

            rows.append(
                {
                    "direction": spec["direction"],
                    "policy_code": policy,
                    "unit_id": unit_id,
                    "final_matching_ids": "; ".join(matches),
                    "final_label": label,
                    "final_reason": reason,
                    "resolution": "adjudicated",
                }
            )

    return pd.DataFrame(rows), errors


def counts_by_policy(frame: pd.DataFrame, direction: str) -> dict:
    subset = frame[frame["direction"] == direction]
    result: dict[str, dict] = {}
    for policy, group in subset.groupby("policy_code", sort=True):
        result[policy] = {
            "total": int(len(group)),
            "labels": {str(k): int(v) for k, v in Counter(group["final_label"]).items()},
        }
    return result


def main() -> None:
    adjudicated, errors = read_adjudicated()
    unanimous = pd.read_csv(UNANIMOUS, dtype=str).fillna("")
    expected_columns = [
        "direction", "policy_code", "unit_id", "final_matching_ids",
        "final_label", "final_reason", "resolution",
    ]
    if list(unanimous.columns) != expected_columns:
        errors.append(f"Unexpected unanimous CSV columns: {list(unanimous.columns)}")
    unanimous["direction"] = unanimous["direction"].replace(
        {"candidate": "candidate_to_reference", "reference": "reference_to_candidate"}
    )

    final = pd.concat([unanimous, adjudicated], ignore_index=True)
    duplicates = final[final.duplicated(["direction", "unit_id"], keep=False)]
    if not duplicates.empty:
        errors.append(f"Duplicate final decisions: {duplicates['unit_id'].tolist()}")

    direction_counts = Counter(final["direction"])
    if direction_counts != {"candidate_to_reference": 114, "reference_to_candidate": 165}:
        errors.append(f"Unexpected direction totals: {dict(direction_counts)}")

    candidate = final[final["direction"] == "candidate_to_reference"]
    reference = final[final["direction"] == "reference_to_candidate"]
    candidate_counts = Counter(candidate["final_label"])
    reference_counts = Counter(reference["final_label"])

    ref_total = len(reference)
    full = reference_counts.get("Fully represented", 0)
    partial = reference_counts.get("Partly represented", 0)
    missing = reference_counts.get("Not represented", 0)
    metrics = {
        "strict_completeness_percent": round(100 * full / ref_total, 2),
        "at_least_partial_percent": round(100 * (full + partial) / ref_total, 2),
        "weighted_completeness_percent": round(100 * (full + 0.5 * partial) / ref_total, 2),
        "not_represented_percent": round(100 * missing / ref_total, 2),
    }

    # Independent-rating agreement before adjudication.
    ratings_path = BASE / "v3_researcher_comparison" / "validation" / "v3_researcher_ratings_all.csv"
    agreement = {}
    if ratings_path.exists():
        ratings = pd.read_csv(ratings_path, dtype=str).fillna("")
        label_col = next((c for c in ratings.columns if c.lower() == "label"), None)
        match_col = next((c for c in ratings.columns if "matching" in c.lower()), None)
        evaluator_col = next(
            (c for c in ratings.columns if c.lower() in {"evaluator", "researcher"}), None
        )
        if label_col and evaluator_col:
            label_wide = ratings.pivot(index=["direction", "unit_id"], columns=evaluator_col, values=label_col)
            label_unanimous = int(label_wide.nunique(axis=1).eq(1).sum())
            agreement["label_unanimous_items"] = label_unanimous
            agreement["label_unanimity_percent"] = round(100 * label_unanimous / len(label_wide), 2)
            if match_col:
                tmp = ratings.copy()
                tmp["_norm_match"] = tmp[match_col].map(lambda x: ";".join(sorted(split_ids(x))))
                match_wide = tmp.pivot(index=["direction", "unit_id"], columns=evaluator_col, values="_norm_match")
                both = label_wide.nunique(axis=1).eq(1) & match_wide.nunique(axis=1).eq(1)
                agreement["label_and_match_unanimous_items"] = int(both.sum())
                agreement["label_and_match_unanimity_percent"] = round(100 * both.mean(), 2)

    result = {
        "status": "valid" if not errors else "invalid",
        "validation_errors": errors,
        "total_final_decisions": int(len(final)),
        "unanimous_decisions": int(len(unanimous)),
        "adjudicated_decisions": int(len(adjudicated)),
        "candidate_units": {
            "total": int(len(candidate)),
            "label_counts": {str(k): int(v) for k, v in sorted(candidate_counts.items())},
            "by_policy": counts_by_policy(final, "candidate_to_reference"),
        },
        "human_reference_units": {
            "total": int(ref_total),
            "label_counts": {str(k): int(v) for k, v in sorted(reference_counts.items())},
            "metrics": metrics,
            "by_policy": counts_by_policy(final, "reference_to_candidate"),
        },
        "independent_researcher_agreement": agreement,
        "earlier_v2_comparison": {
            "gemini_units": 48,
            "human_reference_units": 165,
            "fully_represented": 51,
            "partly_represented": 23,
            "not_represented": 91,
            "strict_completeness_percent": round(100 * 51 / 165, 2),
            "at_least_partial_percent": round(100 * (51 + 23) / 165, 2),
            "weighted_completeness_percent": round(100 * (51 + 0.5 * 23) / 165, 2),
        },
    }

    final = final.sort_values(["direction", "policy_code", "unit_id"]).reset_index(drop=True)
    final.to_csv(OUTPUT_CSV, index=False, encoding="utf-8-sig")
    OUTPUT_JSON.write_text(json.dumps(result, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(result, indent=2, ensure_ascii=False))
    if errors:
        raise SystemExit(1)


if __name__ == "__main__":
    main()
