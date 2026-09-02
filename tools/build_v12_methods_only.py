from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "final_report" / "METHODS_AND_MATERIALS_80_PLUS_DRAFT.docx"
FONT = "Times New Roman"
LIGHT_GRAY = "EFEFEF"


def set_font(run, size=10, bold=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.6)
    section.bottom_margin = Cm(4.6)
    section.left_margin = Cm(4.4)
    section.right_margin = Cm(4.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 12, 12, 7),
        ("Heading 2", 10, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = None
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True

    equation = doc.styles.add_style("Equation", 1)
    equation.base_style = normal
    equation.font.name = FONT
    equation._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    equation._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    equation.font.size = Pt(9.5)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(2)
    equation.paragraph_format.space_after = Pt(4)
    equation.paragraph_format.keep_together = True

    add_page_number(section.footer.paragraphs[0])


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)
    return p


def shade_cell(cell, fill=LIGHT_GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths, indent=80):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, caption, headers, rows, widths, font_size=8.2):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for i, value in enumerate(headers):
        shade_cell(header.cells[i])
        para = header.cells[i].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        set_font(para.add_run(value), font_size, True)
    for values in rows:
        row = table.add_row()
        for i, value in enumerate(values):
            para = row.cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            set_font(para.add_run(str(value)), font_size)
    set_table_geometry(table, widths)
    return table


def count_words(doc):
    pieces = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            pieces.extend(cell.text for cell in row.cells)
    return len(re.findall(r"\b[\w'-]+\b", "\n".join(pieces)))


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    doc.add_heading("3 Methods and Materials", level=1)

    doc.add_heading("3.1 Study Design and Corpus", level=2)
    add_body(
        doc,
        "A controlled within-document factorial design crossed policy, model family, prompting strategy and prompt set. Three policies were each processed under three models, three strategies and two matched prompt sets, producing 54 summaries. Every safety-focused output was paired with a basic output generated from the same policy, model and strategy, giving 27 intervention pairs for RQ2. One output was retained per condition, including valid but poor outputs. The design enables controlled case comparison but, because policies were purposively selected, does not support population-wide claims about policies, models or readers."
    )
    add_table(
        doc,
        "Table 1. Frozen privacy-policy corpus.",
        ["ID", "Organisation", "Role", "Words", "Retrieved"],
        [
            ["PILOT01", "Mozilla", "Short", "857", "12 Aug 2026"],
            ["PILOT02", "DuckDuckGo", "Medium", "2,077", "12 Aug 2026"],
            ["PILOT03", "Automattic", "Long/complex", "8,285", "12 Aug 2026"],
        ],
        [1500, 1900, 1600, 1100, 1900],
    )
    add_body(
        doc,
        "The corpus comprised publicly accessible, contemporary English-language policies representing short, medium and long/complex documents. Complete webpages, retrieval metadata and cleaned experimental text were stored separately. Cleaning removed navigation, duplicated webpage furniture and malformed spacing while preserving substantive wording, headings, conditions, qualifications and exceptions. Word counts and SHA-256 hashes were recorded, and the identical frozen text was supplied to every generator and evaluator. This reduced irrelevant webpage noise without altering the evidence base used for verification."
    )

    doc.add_heading("3.2 Models, Prompt Intervention and Generation", level=2)
    add_table(
        doc,
        "Table 2. Frozen model endpoints.",
        ["Family", "Provider", "Model identifier"],
        [
            ["GPT", "OpenAI", "gpt-5.6-luna"],
            ["Llama", "Together AI", "meta-llama/Llama-3.3-70B-Instruct-Turbo"],
            ["Mistral", "Mistral API", "mistral-large-2512"],
        ],
        [1250, 1550, 5200],
    )
    add_body(
        doc,
        "The systems represent three model families and provider ecosystems available through comparable APIs. Selection was based on practical access, capacity to process complete policies and diversity of development contexts; it is not treated as a controlled open-versus-closed comparison because training data, architecture, scale and deployment also differ. Direct prompting stated the task and audience with minimal guidance. Role-guided prompting added a plain-language communication role. Structured prompting instructed the model to identify and organise important material internally before returning only the summary; it did not request disclosure of hidden reasoning."
    )
    add_body(
        doc,
        "The basic prompt requested a clear, concise plain-English summary for a general adult reader. Its matched safety-focused version retained the task, audience and prose format but required source-only information, preservation of conditions and qualifications, avoidance of assumptions and retention of information important to the reader. This controlled change isolates the safety intervention more credibly than changing task and format simultaneously. Full frozen templates are reported in Appendix C."
    )
    add_body(
        doc,
        "Llama and Mistral used temperature 0.0. The GPT endpoint rejected the temperature parameter, so it was omitted and logged as a provider constraint. The maximum output allowance was 6,000 tokens. Logs preserved exact model and prompt versions, timestamps, token usage, response identifiers, status and errors. Successful outputs were not regenerated for quality reasons. Model and prompt identities were replaced by blind summary IDs before evaluation."
    )

    doc.add_heading("3.3 Evaluation Framework", level=2)
    add_body(
        doc,
        "RQ1 combined accessibility, conciseness and source alignment. Readability was estimated per summary using Flesch Reading Ease, Flesch-Kincaid Grade and SMOG Grade. Conciseness used word count and compression ratio:"
    )
    doc.add_paragraph("Compression ratio = summary words / source-policy words", style="Equation")
    add_body(
        doc,
        "Higher Reading Ease and lower grade estimates conventionally indicate simpler textual structure, while a lower compression ratio indicates a shorter summary relative to its source. These formulae do not measure comprehension, legal adequacy or usability and may be affected by sentence segmentation and formatting."
    )
    add_body(
        doc,
        "Summary-to-source verification checked hallucination and distortion. Summaries were deterministically divided into visible sentences and assessed against the complete frozen policy. Gemini 3.5 Flash Lite assigned Supported, Partially supported, Contradicted or Unsupported and returned a failure type, exact evidence and reason. Partial support normally operationalised distortion; unsupported content normally operationalised hallucination. MiniCheck with FLAN-T5-Large independently returned a binary support prediction and probability for the identical sentence IDs. MiniCheck was not converted into Gemini's four categories and was not used to measure omission."
    )
    add_body(
        doc,
        "Coverage used the reverse direction. Gemini first produced a preliminary policy-specific set of information important to a non-specialist reader and scored each item against every associated summary as Covered, Partially covered or Not covered. Final omission analysis uses independently proposed researcher units: proposals must be source-supported, deduplicated and adjudicated into one shared denominator before researchers and Gemini score identical unit-summary pairs. A topic absent from the policy therefore cannot be counted as omitted. Strict coverage equals Covered divided by all units; weighted coverage equals (Covered + 0.5 x Partially covered) divided by all units. The 0.5 weight is a transparent convention rather than a natural ground truth."
    )

    doc.add_heading("3.4 Researcher Validation, Analysis and Prototype", level=2)
    add_body(
        doc,
        "Three student researchers receive identical blinded packages containing the complete policies, anonymised summaries and frozen sentence IDs without model, prompt or automated labels. For each sentence they independently record a label, failure type, exact source evidence, reason and confidence. For coverage they first identify important information from the source without viewing summaries, then score the adjudicated units against the summaries. Independent work precedes adjudication to limit influence between researchers, although the evaluators are project members rather than representative end users or legal experts."
    )
    add_body(
        doc,
        "Researcher reliability will be reported using label distributions, pairwise exact agreement, pairwise Cohen's kappa and confusion matrices. Automated-researcher comparisons use identical sentence and coverage IDs, with both four-class and binary faithfulness views where appropriate. Class prevalence will accompany agreement coefficients because a high value may be dominated by Supported cases. RQ2 uses safety-focused-minus-basic differences within each matched pair and reports mean and median differences plus positive, tied and negative pair counts. With only three policies, interpretation emphasises magnitude, direction and consistency rather than population-level significance."
    )
    add_body(
        doc,
        "The PolicyLens Streamlit prototype implements evidence-linked inspection. A user selects policy, model, strategy and prompt set, then explores the summary, sentence labels, failure types, exact evidence, reasons, MiniCheck results, coverage units and evaluator disagreements. Consistent colour coding distinguishes supported or covered, partial or distorted, unsupported or omitted, and pending states. The prototype demonstrates research transparency; it is neither a legal-advice service nor evidence of deployment safety."
    )

    words = count_words(doc)
    doc.core_properties.title = "Methods and Materials - Verifying LLM-Generated Plain-English Summaries of Privacy Policies"
    doc.core_properties.subject = "Standalone MSc AI dissertation methods draft"
    doc.core_properties.author = "MSc AI Project Team"
    doc.core_properties.keywords = "LLM, privacy policy, matched intervention, faithfulness, coverage"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"WORDS={words}")


if __name__ == "__main__":
    build()
