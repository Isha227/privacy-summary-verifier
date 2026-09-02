from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "final_report" / "Report_Structure_and_Marking_Guide.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "667085"
PALE_BLUE = "E8EEF5"
PALE_GRAY = "F2F4F7"
PALE_GOLD = "FFF4CC"
WHITE = "FFFFFF"
BLACK = "111827"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = tr_pr.find(qn("w:cantSplit"))
    if cant_split is None:
        cant_split = OxmlElement("w:cantSplit")
        tr_pr.append(cant_split)


def set_table_geometry(table, widths_dxa, indent_dxa=120):
    total = sum(widths_dxa)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for i, (cell, width) in enumerate(zip(row.cells, widths_dxa)):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_bullet(doc, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_together = True
    p.add_run(text)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.keep_together = True
    p.add_run(text)
    return p


def add_label_paragraph(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    r = p.add_run(label)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p.add_run(text)
    return p


def add_callout(doc, title, text, fill=PALE_BLUE):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    set_row_cant_split(table.rows[0])
    set_table_geometry(table, [9360], indent_dxa=120)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.keep_together = True
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=NAVY, bold=True)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.keep_together = True
    p2.paragraph_format.space_after = Pt(0)
    return table


def add_two_col_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    set_repeat_table_header(header)
    for i, text in enumerate(headers):
        set_cell_shading(header.cells[i], PALE_BLUE)
        p = header.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        set_run_font(r, size=10, color=NAVY, bold=True)
    for row_data in rows:
        row = table.add_row()
        set_row_cant_split(row)
        cells = row.cells
        for i, text in enumerate(row_data):
            p = cells[i].paragraphs[0]
            p.paragraph_format.keep_together = True
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(str(text))
            set_run_font(r, size=9.5, color=BLACK)
            if i == 1 and str(text).replace(",", "").isdigit():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_table_geometry(table, widths)
    return table


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    settings = {
        "Title": (26, NAVY, 0, 8),
        "Subtitle": (13, MUTED, 0, 16),
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in settings.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Bullet 2", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25
    styles["List Bullet"].paragraph_format.left_indent = Inches(0.375)
    styles["List Bullet"].paragraph_format.first_line_indent = Inches(-0.188)
    styles["List Bullet 2"].paragraph_format.left_indent = Inches(0.65)
    styles["List Bullet 2"].paragraph_format.first_line_indent = Inches(-0.188)
    styles["List Number"].paragraph_format.left_indent = Inches(0.375)
    styles["List Number"].paragraph_format.first_line_indent = Inches(-0.188)


def add_section_plan(doc, number, title, words, purpose, content, marker_focus, required_output):
    doc.add_heading(f"{number}. {title}", level=1)
    add_label_paragraph(doc, "Word budget: ", words)
    add_label_paragraph(doc, "Purpose: ", purpose)
    doc.add_heading("What to include", level=2)
    for item in content:
        add_bullet(doc, item)
    doc.add_heading("What the marker should be able to see", level=2)
    for item in marker_focus:
        add_bullet(doc, item)
    add_callout(doc, "Section output", required_output, PALE_GRAY)


def build_document():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    configure_styles(doc)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hr = header.add_run("MSc AI Dissertation | Report Planning Guide")
    set_run_font(hr, size=8.5, color=MUTED)
    add_page_number(section.footer.paragraphs[0])

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kr = kicker.add_run("REPORT STRUCTURE AND MARKING PLAN")
    set_run_font(kr, size=10, color=BLUE, bold=True)
    kicker.paragraph_format.space_before = Pt(42)
    kicker.paragraph_format.space_after = Pt(10)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Verifying LLM-Generated Plain-English Summaries of Privacy Policies")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("A practical blueprint for a 4,000-word LNCS-style dissertation report")

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.space_before = Pt(10)
    lead.paragraph_format.space_after = Pt(18)
    run = lead.add_run(
        "Central argument: LLMs may make privacy policies more accessible, but their summaries must be verified for hallucination, distortion and omission."
    )
    set_run_font(run, size=12, color=NAVY, bold=True)

    add_callout(
        doc,
        "Conference-paper angle",
        "The comparison of three models and three prompting strategies is the experiment. The wider purpose is to determine whether plain-English privacy-policy summaries can be made accessible while remaining aligned with the source, and whether their reliability can be checked using automated and researcher evaluation.",
        PALE_GOLD,
    )

    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    doc.add_heading("1. The report's narrative", level=1)
    narrative = [
        "Privacy policies are difficult for ordinary readers to navigate.",
        "LLMs may improve accessibility through summarisation and simplification.",
        "Simplification can introduce hallucination, distortion and omission.",
        "Reliable evaluation therefore requires both summary-to-source and source-to-summary verification.",
        "The study tests generation quality, a safety-focused prompting intervention and evaluator agreement.",
        "The prototype exposes the evidence behind each judgement so that results can be inspected rather than merely accepted.",
    ]
    for step in narrative:
        add_number(doc, step)

    doc.add_heading("2. Word budget", level=1)
    add_two_col_table(
        doc,
        ["Report section", "Suggested words", "Primary assessment contribution"],
        [
            ("Introduction", "550", "Rationale, scope, contribution and research questions"),
            ("Relevant Work", "800", "Critical literature synthesis and research gap"),
            ("Methods & Materials", "1,050", "Proposed approach and reproducibility"),
            ("Experimental Results and Discussion", "1,300", "Evidence, interpretation and RQ answers"),
            ("Conclusions", "300", "Overall findings, limitations and future work"),
            ("Total", "4,000", "Excludes title, abstract and references"),
        ],
        [2700, 1500, 5160],
    )
    note = doc.add_paragraph()
    nr = note.add_run("Important: ")
    nr.bold = True
    note.add_run(
        "until the university confirms otherwise, treat appendix text as potentially countable. Keep all essential reasoning and methodology in the main report."
    )

    doc.add_heading("3. Front matter", level=1)
    doc.add_heading("Title and author details", level=2)
    add_bullet(doc, "Use the approved title exactly as shown on the cover of this guide.")
    add_bullet(doc, "Include all author names and UWE student IDs.")
    doc.add_heading("Abstract (100-150 words)", level=2)
    add_bullet(doc, "State the problem, purpose, design, evaluation methods, principal findings and conclusion.")
    add_bullet(doc, "Avoid citations, unexplained abbreviations and implementation detail.")
    doc.add_heading("Keywords", level=2)
    doc.add_paragraph(
        "Privacy Policies; Large Language Models; Plain-English Summarisation; Factual Consistency; Human Evaluation"
    )

    add_section_plan(
        doc,
        "4",
        "Introduction",
        "Approximately 550 words",
        "Establish why the problem matters, define the controlled scope and present the questions that guide the study.",
        [
            "Background: privacy policies affect user decisions but are often long, technical and difficult to understand.",
            "Motivation: LLMs may improve access, but fluent language is not evidence of correctness or completeness.",
            "Central tension: accessibility must be balanced against information preservation.",
            "Scope: three policies, three models, three prompt strategies, two prompt sets and 54 summaries.",
            "Contribution: bidirectional verification, automated-researcher comparison and an evidence-linked prototype.",
            "The three research questions, followed by one short paragraph describing the remaining report sections.",
        ],
        [
            "A convincing problem linked to ordinary users and wider legal-information access.",
            "A feasible scope rather than an unsupported claim about all legal documents or all LLMs.",
            "Clear, sound research questions that are answerable using the reported experiment.",
        ],
        "A focused opening that frames verification as the project purpose and the model/prompt comparison as the experimental design.",
    )

    doc.add_heading("Approved research questions", level=2)
    for label, rq in [
        ("RQ1", "To what extent can the selected LLMs generate privacy-policy summaries that are readable, concise and aligned with the source?"),
        ("RQ2", "How effectively does safety-focused prompting reduce hallucination, distortion and omission while maintaining readability and conciseness?"),
        ("RQ3", "To what extent do automated evaluator judgements align with researcher judgements when identifying hallucination, distortion and omission?"),
    ]:
        add_label_paragraph(doc, f"{label}: ", rq)

    add_section_plan(
        doc,
        "5",
        "Relevant Work",
        "Approximately 800 words",
        "Build the intellectual justification for the study and identify a precise gap rather than producing a catalogue of papers.",
        [
            "Legal language and plain English: jargon, syntax, accessibility and the distinction between readability and understanding.",
            "Privacy policies: length, complexity, user engagement and previous simplification or summarisation work.",
            "LLM summarisation and prompting: abstractive summarisation, simplification, model families and prompt design.",
            "Reliability failures: hallucination, factual inconsistency, distortion of qualifications and omission of important information.",
            "Evaluation methods: readability, compression, faithfulness, coverage, automated evaluators and human agreement.",
            "Gap: accessible summaries require bidirectional verification, while automated evaluator validity should be checked against researcher judgement.",
        ],
        [
            "Critical comparison across studies, including limitations and contradictions.",
            "Recent state-of-the-art work alongside necessary foundational sources.",
            "A visible connection between the literature gap, the RQs and every evaluation choice.",
        ],
        "A thematic argument that leads naturally to the proposed method. Each paragraph should compare evidence, identify a limitation or justify a project decision.",
    )

    add_section_plan(
        doc,
        "6",
        "Methods & Materials",
        "Approximately 1,050 words",
        "Describe the study clearly enough that another researcher could reproduce it and understand every methodological choice.",
        [
            "Study design: a controlled, predominantly quantitative mixed-methods case study with 27 matched basic-safety pairs.",
            "Corpus: Mozilla, DuckDuckGo and Automattic policies selected as short, medium and long/complex cases; include collection, cleaning and freezing procedures.",
            "Models: GPT-5.6 Luna, Llama 3.3 70B Instruct Turbo and Mistral Large 3, with providers and a cautious justification of diversity.",
            "Prompting conditions: direct, role-guided and structured strategies under basic and safety-focused prompt sets.",
            "Generation: API configuration, logging, token records, error handling and preservation of valid-but-poor outputs.",
            "Accessibility and conciseness: Flesch Reading Ease, Flesch-Kincaid Grade, SMOG, word count and compression ratio.",
            "Summary-to-source verification: Supported, Partially supported/distortion, Unsupported/hallucination and Contradicted where applicable.",
            "Source-to-summary coverage: Covered, Partially covered and Not covered/omission, using source information actually present and judged important.",
            "Automated evaluation: MiniCheck sentence scoring and Gemini evidence-based assessment.",
            "Researcher evaluation: three blinded independent annotators, evidence and reasons, source-unit adjudication and final shared coverage denominator.",
            "Prototype: selectors, colour-coded labels, sentence-level evidence, coverage view and disagreement explorer.",
        ],
        [
            "A strong, clearly explained proposed approach rather than a list of software tools.",
            "Justification for each model, prompt strategy, measure and evaluator.",
            "Transparency about design constraints, human roles and reproducibility controls.",
        ],
        "A concise method that supports all three RQs. Full prompts, lengthy instructions and supplementary technical details should be moved to appendices, not omitted.",
    )

    doc.add_heading("Experimental matrix", level=2)
    add_callout(
        doc,
        "Controlled design",
        "3 policies x 3 models x 3 prompt strategies x 2 prompt sets = 54 summaries. The basic and safety-focused conditions create 27 matched pairs.",
        PALE_BLUE,
    )

    doc.add_heading("Bidirectional verification", level=2)
    add_label_paragraph(
        doc,
        "Summary-to-source: ",
        "each generated sentence is checked against the policy to identify supported content, distortion and hallucination.",
    )
    add_label_paragraph(
        doc,
        "Source-to-summary: ",
        "important information found in the source is checked against the summary to identify full coverage, partial coverage and omission.",
    )

    add_section_plan(
        doc,
        "7",
        "Experimental Results and Discussion",
        "Approximately 1,300 words",
        "Present verified evidence and interpret it by research question. This is the largest section because it carries the empirical argument.",
        [
            "Completion and data-quality checks: 54 successful summaries, 27 matched pairs, evaluator coverage and researcher validation status.",
            "RQ1 results: baseline readability, conciseness and source alignment by model and prompt strategy.",
            "RQ2 results: paired safety-minus-basic changes in readability, length, faithfulness and coverage.",
            "RQ3 results: automated-researcher agreement, Cohen's kappa where suitable, agreement by failure type and disagreement cases.",
            "Qualitative failure analysis: one clear hallucination, one qualification distortion, one omission and one evaluator disagreement, each supported by evidence.",
            "Prototype demonstration: one figure showing how the interface exposes labels, evidence and reasons.",
            "Overall discussion: relate findings to the literature and explain the accessibility-reliability trade-off.",
            "Limitations: three policies, three model families, changing APIs, formula limitations, researcher subjectivity, project-team annotators and no independent end-user comprehension study.",
        ],
        [
            "Clear tables or figures rather than unstructured numerical reporting.",
            "Direct answers to each RQ, including results that do not support expectations.",
            "Critical interpretation, comparison with prior research and appropriately cautious conclusions.",
        ],
        "An RQ-led results narrative. Do not organise this section around CSV filenames, APIs or the order in which scripts were run.",
    )

    doc.add_heading("Current automated pattern to report cautiously", level=2)
    add_bullet(doc, "Safety-focused summaries were approximately 182 words longer on average.")
    add_bullet(doc, "Gemini strict support increased by approximately 2.4 percentage points.")
    add_bullet(doc, "Preliminary strict coverage increased by approximately 15.3 percentage points.")
    add_bullet(doc, "Readability generally became slightly harder.")
    add_bullet(doc, "MiniCheck support did not improve consistently.")
    add_callout(
        doc,
        "Interpretation",
        "The safety-focused prompt produced a trade-off rather than a universal improvement: it retained more important information, but usually produced longer and somewhat less readable summaries. Final claims must incorporate the completed researcher results and the shared human-adjudicated coverage units.",
        PALE_GOLD,
    )

    add_section_plan(
        doc,
        "8",
        "Conclusions",
        "Approximately 300 words",
        "Close the argument by answering the RQs, stating the overall contribution and identifying realistic next steps.",
        [
            "Give a one- or two-sentence answer to each RQ without introducing new results.",
            "State the overall conclusion that accessibility alone is insufficient for legal-information summaries.",
            "Explain that verification must examine what a summary says, whether meaning changed and what important information is missing.",
            "Propose future work: a larger corpus, independent experts, end-user comprehension, improved evidence retrieval, correction mechanisms and other languages or legal-document types.",
        ],
        [
            "Conclusions that are supported by the experiment and acknowledge the small case-study scope.",
            "A clear contribution and feasible next steps, rather than a repetition of the abstract.",
        ],
        "A concise ending that tells the reader what was learned, why it matters and what should happen next.",
    )

    doc.add_heading("9. Appendices", level=1)
    add_label_paragraph(doc, "Appendix A: ", "GitHub repository link and reproducibility information.")
    add_label_paragraph(doc, "Appendix B: ", "System and evaluation diagrams with numbered titles and citations from the main text.")
    add_label_paragraph(doc, "Appendix C, if permitted: ", "Complete frozen prompt templates.")
    add_label_paragraph(doc, "Appendix D, if permitted: ", "Researcher instructions, label definitions and supplementary tables.")
    add_callout(
        doc,
        "Appendix rule",
        "Do not hide information needed to understand or reproduce the central method in an appendix. Confirm whether additional appendices are allowed and whether appendix words count toward the 4,000-word limit.",
        PALE_GOLD,
    )

    doc.add_heading("10. Marking-criteria map", level=1)
    add_two_col_table(
        doc,
        ["Criterion", "Weight", "Where the report earns the marks"],
        [
            ("Research rationale, scope and questions", "25%", "Introduction: wider context, feasible scope, novelty and precise RQs"),
            ("Literature review", "20%", "Relevant Work: state of the art, critical synthesis and defensible gap"),
            ("Proposed approach", "25%", "Methods: clear design, justifications, reproducibility and prototype"),
            ("Results, discussion and conclusions", "20%", "RQ-led evidence, meaningful interpretation and supported conclusions"),
            ("Report structure", "10%", "LNCS format, coherent flow, citations, labelled figures/tables and word control"),
        ],
        [2900, 900, 5560],
    )

    doc.add_heading("11. Final writing checks", level=1)
    checks = [
        "Every research question is introduced, operationalised in Methods and answered in Results and Conclusions.",
        "Every metric is justified by literature and interpreted according to what it actually measures.",
        "Readability is not described as proof of comprehension.",
        "Preliminary Gemini coverage is not confused with final human-adjudicated coverage.",
        "Automated-automated agreement is not presented as validation; RQ3 requires comparison with researchers.",
        "Claims remain proportionate to a three-policy controlled case study.",
        "All tables and figures are numbered, titled and discussed in the text.",
        "The main report follows LNCS formatting and stays within the verified word-count rules.",
        "References are complete and consistently formatted.",
        "The final paragraph states the project's practical message clearly.",
    ]
    for item in checks:
        add_bullet(doc, item)

    add_callout(
        doc,
        "Final practical message",
        "LLMs can produce useful plain-English privacy-policy summaries, but a summary should not be trusted merely because it is fluent and readable. It must be checked for unsupported statements, altered meaning and important missing information.",
        PALE_BLUE,
    )

    doc.core_properties.title = "Report Structure and Marking Guide"
    doc.core_properties.subject = "MSc AI dissertation planning"
    doc.core_properties.author = "MSc AI Project Team"
    doc.core_properties.keywords = "LNCS, MSc AI, privacy policies, LLM summarisation, marking criteria"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
