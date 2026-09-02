from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "final_report" / "CONCLUSIONS_80_PLUS_DRAFT.docx"
FONT = "Times New Roman"
RED = "A00000"


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.6)
    section.bottom_margin = Cm(4.6)
    section.left_margin = Cm(4.4)
    section.right_margin = Cm(4.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True

    heading = doc.styles["Heading 1"]
    heading.font.name = FONT
    heading._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    heading._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    heading.font.size = Pt(12)
    heading.font.bold = True
    heading.font.color.rgb = None
    heading.paragraph_format.space_before = Pt(12)
    heading.paragraph_format.space_after = Pt(7)
    heading.paragraph_format.keep_with_next = True

    pending = doc.styles.add_style("Pending Result", 1)
    pending.base_style = normal
    pending.font.name = FONT
    pending._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    pending._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    pending.font.size = Pt(9.5)
    pending.font.bold = True
    pending.font.color.rgb = RGBColor.from_string(RED)
    pending.paragraph_format.space_before = Pt(4)
    pending.paragraph_format.space_after = Pt(5)
    pending.paragraph_format.keep_together = True

    add_page_number(section.footer.paragraphs[0])


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)


def count_words(doc):
    return len(re.findall(r"\b[\w'-]+\b", "\n".join(p.text for p in doc.paragraphs)))


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    doc.add_heading("5 Conclusions", level=1)

    add_body(
        doc,
        "This study examined whether large language models can transform privacy policies into plain-English summaries that are concise, readable and aligned with their sources. Across three policies, three model families and three prompting strategies, 54 summaries formed 27 matched comparisons between a basic prompt and a safety-focused version. The completed automated analysis shows that the task cannot be judged through fluency or readability alone: summaries may be shorter and apparently clear while still omitting information, weakening qualifications or receiving different assessments from different evaluators."
    )
    add_body(
        doc,
        "The models produced distinct trade-offs. GPT generated the longest summaries and achieved the highest preliminary coverage, Llama was the most concise, and Mistral obtained the most favourable formula-based readability values. Nevertheless, all model-level mean grade estimates remained above 12, so none can be described as consistently accessible to every ordinary reader. Safety-focused prompting increased preliminary strict coverage by 15.29 percentage points and weighted coverage by 13.02 points across the matched pairs. It also produced summaries that were approximately 182 words longer and generally harder according to the readability formulae. Its effect on faithfulness was evaluator-dependent: Gemini support increased, whereas MiniCheck support decreased slightly. The intervention therefore improved information retention but did not create an unqualified improvement across all quality dimensions."
    )
    pending = doc.add_paragraph(style="Pending Result")
    pending.add_run(
        "[PENDING FINAL VALIDATION] Insert the final RQ3 conclusion after the A1/A2/A3 researcher judgements, adjudication and automated-researcher agreement analyses have been completed."
    )
    add_body(
        doc,
        "The project contributes a practical verification-oriented prototype and a bidirectional evaluation design. Summary-to-source assessment identifies unsupported or meaning-altering statements, while source-to-summary assessment identifies important information that has been omitted. This distinction is important because a summary can contain only supported sentences yet remain incomplete. Evidence-linked labels and reasons also make individual judgements inspectable rather than reducing reliability to a single opaque score."
    )
    add_body(
        doc,
        "The findings remain limited to three purposively selected policies, one generation per condition, formula-based readability estimates and changing external API systems. The researchers are project members rather than representative end users or legal professionals. Future work should therefore extend the study to a broader policy sample, test whether ordinary readers actually understand and use the summaries, obtain legal-expert review of consequential distortions and omissions, and examine whether adaptive length or document-specific prompting can improve coverage without sacrificing accessibility. Subject to the pending human validation, the central conclusion is that LLMs can assist privacy-policy communication, but their outputs require transparent, source-grounded verification before they should be relied upon as accurate legal information."
    )

    doc.core_properties.title = "Conclusions - Verifying LLM-Generated Plain-English Summaries of Privacy Policies"
    doc.core_properties.subject = "Standalone MSc AI dissertation conclusion draft"
    doc.core_properties.author = "MSc AI project team"
    doc.core_properties.keywords = "privacy policy, summarisation, LLM, verification, conclusion"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"WORDS={count_words(doc)}")


if __name__ == "__main__":
    build()
