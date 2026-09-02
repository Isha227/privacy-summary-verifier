from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "final_report" / "EXPERIMENTAL_RESULTS_AND_DISCUSSION_80_PLUS_DRAFT.docx"
ASSET = ROOT / "outputs" / "final_report" / "v12_draft_assets" / "figure_results_pair_directions.png"
FONT = "Times New Roman"
LIGHT_GRAY = "EFEFEF"
RED = "A00000"


def set_font(run, size=10, bold=False, color=None):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.6)
    section.bottom_margin = Cm(4.6)
    section.left_margin = Cm(4.4)
    section.right_margin = Cm(4.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.5)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True
    for name, size, before, after in (("Heading 1", 12, 12, 7), ("Heading 2", 10, 8, 4)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
    caption = doc.styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = None
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(5)
    caption.paragraph_format.keep_with_next = True
    placeholder = doc.styles.add_style("Pending Result", 1)
    placeholder.base_style = normal
    placeholder.font.name = FONT
    placeholder._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    placeholder._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    placeholder.font.size = Pt(9.5)
    placeholder.font.bold = True
    placeholder.font.color.rgb = RGBColor.from_string(RED)
    placeholder.paragraph_format.space_before = Pt(4)
    placeholder.paragraph_format.space_after = Pt(5)
    placeholder.paragraph_format.keep_together = True
    add_page_number(section.footer.paragraphs[0])


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)
    return p


def shade_cell(cell, fill=LIGHT_GRAY):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_table_geometry(table, widths, indent=80):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        if tr_pr.find(qn("w:cantSplit")) is None:
            tr_pr.append(OxmlElement("w:cantSplit"))
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def add_table(doc, caption, headers, rows, widths, numeric_cols=None, font_size=8.0):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    tr_pr = header._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    tr_pr.append(repeat)
    for i, value in enumerate(headers):
        shade_cell(header.cells[i])
        para = header.cells[i].paragraphs[0]
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.paragraph_format.space_after = Pt(0)
        set_font(para.add_run(value), font_size, True)
    for values in rows:
        row = table.add_row()
        for i, value in enumerate(values):
            para = row.cells[i].paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            if numeric_cols and i in numeric_cols:
                para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(para.add_run(str(value)), font_size)
    set_table_geometry(table, widths)
    return table


def load_font(size, bold=False):
    path = Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf")
    return ImageFont.truetype(str(path), size) if path.exists() else ImageFont.load_default()


def make_direction_figure():
    ASSET.parent.mkdir(parents=True, exist_ok=True)
    rows = [
        ("Word count", 24, 0, 3),
        ("Flesch Reading Ease", 8, 0, 19),
        ("Gemini support", 16, 4, 7),
        ("MiniCheck support", 10, 3, 14),
        ("Strict coverage", 22, 4, 1),
        ("Weighted coverage", 26, 0, 1),
    ]
    width, height = 1500, 680
    img = Image.new("RGB", (width, height), "#FFFFFF")
    draw = ImageDraw.Draw(img)
    title_font, label_font, value_font = load_font(38, True), load_font(25), load_font(22, True)
    draw.text((65, 35), "Direction of safety-focused minus basic change (27 pairs)", fill="#244A73", font=title_font)
    draw.text((900, 100), "Higher", fill="#2F6B9A", font=value_font)
    draw.text((1050, 100), "Same", fill="#777777", font=value_font)
    draw.text((1180, 100), "Lower", fill="#C56A2D", font=value_font)
    x0, bar_w, y0, row_h = 500, 840, 155, 78
    colours = ["#4F81BD", "#B7B7B7", "#F4A261"]
    for idx, (label, high, same, low) in enumerate(rows):
        y = y0 + idx * row_h
        draw.text((65, y + 9), label, fill="#222222", font=label_font)
        x = x0
        for value, colour in zip((high, same, low), colours):
            w = bar_w * value / 27
            if value:
                draw.rectangle((x, y, x + w, y + 48), fill=colour)
                txt = str(value)
                bbox = draw.textbbox((0, 0), txt, font=value_font)
                if w > 35:
                    draw.text((x + (w - (bbox[2] - bbox[0])) / 2, y + 10), txt, fill="#FFFFFF", font=value_font)
            x += w
        draw.rectangle((x0, y, x0 + bar_w, y + 48), outline="#666666", width=2)
    draw.text((65, 630), "Higher and lower describe metric direction, not universal improvement or deterioration.", fill="#555555", font=load_font(21))
    img.save(ASSET, dpi=(180, 180))


def add_figure(doc):
    make_direction_figure()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ASSET), width=Cm(12.0))
    shape._inline.docPr.set("title", "Figure 1")
    shape._inline.docPr.set("descr", "Stacked bars showing the number of matched pairs with higher, unchanged or lower values under safety-focused prompting for six outcomes")
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run("Fig. 1. Direction of change across 27 matched intervention pairs.")


def count_words(doc):
    pieces = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            pieces.extend(cell.text for cell in row.cells)
    return len(re.findall(r"\b[\w'-]+\b", "\n".join(pieces)))


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)
    doc.add_heading("4 Experimental Results and Discussion", level=1)
    add_body(
        doc,
        "All 54 planned summaries were generated, forming 27 complete matched basic-safety pairs. Gemini and MiniCheck each completed assessment of all 1,393 frozen summary sentences. Gemini labelled 1,345 sentences Supported, 34 Partially supported, 13 Unsupported and one Contradicted; its failure field contained 34 distortions, 13 hallucinations and one combined case. MiniCheck classified 1,255 sentences as supported and 138 as unsupported. Gemini also completed 936 preliminary coverage comparisons: 550 Covered, 192 Partially covered and 194 Not covered. These distributions show why sentence faithfulness cannot establish completeness and why the evaluators should not be treated as interchangeable. Researcher validation remains pending and is not inferred from the automated results."
    )

    doc.add_heading("4.1 RQ1: Summary Quality and Source Alignment", level=2)
    add_table(
        doc,
        "Table 3. Mean automated outcomes by model across 18 summaries.",
        ["Model", "Words", "FRE", "FK", "Comp.", "Gem. sup.", "Mini. sup.", "Prelim. cov."],
        [
            ["GPT", "835.4", "30.88", "14.36", "0.344", "97.89%", "90.73%", "84.77%"],
            ["Llama", "352.2", "36.30", "14.22", "0.187", "94.70%", "93.06%", "34.39%"],
            ["Mistral", "527.3", "43.85", "12.57", "0.260", "94.10%", "90.08%", "50.48%"],
        ],
        [1000, 800, 700, 700, 780, 1050, 1050, 1200],
        numeric_cols={1, 2, 3, 4, 5, 6, 7},
        font_size=7.6,
    )
    add_body(
        doc,
        "The models displayed different quality profiles rather than one universally superior result. GPT produced the longest summaries and the highest preliminary strict coverage, whereas Llama produced the shortest summaries and the lowest preliminary coverage. Mistral achieved the highest Reading Ease and lowest Flesch-Kincaid Grade. Nevertheless, mean Reading Ease remained below 44 and grade estimates above 12 for every model, so the formulae do not support a strong claim of uniformly simple language. They also do not demonstrate actual comprehension."
    )
    add_body(
        doc,
        "The apparent length-coverage relationship is substantial: GPT retained about 34.4% of source words and covered 84.8% of preliminary units, while Llama retained 18.7% and covered 34.4%. This pattern suggests that greater conciseness can remove information a reader may need, although three purposively selected policies cannot establish a general model effect. Gemini gave GPT the highest mean support, but MiniCheck favoured Llama. RQ1 is therefore answered cautiously: the systems produced summaries considerably shorter than their policies and most sentences received automated support, but accessibility remained demanding, coverage varied markedly and source-alignment estimates depended on the evaluator."
    )

    doc.add_heading("4.2 RQ2: Effect of Safety-Focused Prompting", level=2)
    add_table(
        doc,
        "Table 4. Safety-focused minus basic differences across 27 matched pairs.",
        ["Outcome", "Basic", "Safety", "Mean delta", "+ / 0 / -"],
        [
            ["Words", "480.9", "662.4", "+181.56", "24 / 0 / 3"],
            ["Flesch Reading Ease", "38.11", "35.90", "-2.21", "8 / 0 / 19"],
            ["Flesch-Kincaid Grade", "13.31", "14.12", "+0.81", "20 / 0 / 7"],
            ["Gemini strict support", "94.38%", "96.75%", "+2.37 pp", "16 / 4 / 7"],
            ["MiniCheck support", "91.84%", "90.74%", "-1.09 pp", "10 / 3 / 14"],
            ["Prelim. strict coverage", "48.90%", "64.19%", "+15.29 pp", "22 / 4 / 1"],
            ["Prelim. weighted coverage", "61.18%", "74.20%", "+13.02 pp", "26 / 0 / 1"],
        ],
        [2500, 1200, 1200, 1450, 1450],
        numeric_cols={1, 2, 3, 4},
        font_size=7.8,
    )
    add_body(
        doc,
        "The clearest automated intervention effect was broader information retention. Preliminary strict coverage rose by 15.29 percentage points and weighted coverage by 13.02 points; weighted coverage increased in 26 of 27 pairs. The safeguards appear to have encouraged models to preserve more source material. This benefit was not free: safety-focused summaries were 181.6 words longer on average and longer in 24 pairs, while Reading Ease fell in 19 pairs and grade level increased. The result indicates an accessibility-reliability trade-off rather than a universal improvement."
    )
    add_body(
        doc,
        "Faithfulness conclusions were evaluator-dependent. Gemini strict support increased by 2.37 percentage points, but MiniCheck support decreased by 1.09 points and fell in 14 pairs. The difference may reflect Gemini's multi-class, context-rich judgement and MiniCheck's binary document-grounded decision, but the automated data alone cannot establish which interpretation is more credible. RQ2 can therefore be answered provisionally only: safety-focused prompting strongly improved preliminary coverage, but made outputs longer and harder according to formula-based measures, and did not consistently improve sentence support across evaluators."
    )

    doc.add_heading("4.3 RQ3: Automated-Researcher Agreement", level=2)
    pending = doc.add_paragraph(style="Pending Result")
    pending.add_run("[PENDING RESEARCHER RESULTS] Insert validated A1/A2/A3 label distributions, pairwise exact agreement and Cohen's kappa, adjudicated labels, Gemini-human four-class and binary comparisons, MiniCheck-human binary comparison, and shared-unit coverage agreement.")
    add_body(
        doc,
        "The completed Gemini-MiniCheck comparison is a diagnostic, not human validation. RQ3 requires identical sentence and coverage IDs to be compared with independent researcher judgements. Agreement coefficients must be reported with class distributions and confusion matrices because the dominant Supported class could inflate an overall value. Discussion should identify whether disagreements concern subtle changes of scope and qualification, absent evidence or genuinely contradictory content. Until these analyses are complete, no conclusion is made about automated-evaluator reliability."
    )
    pending2 = doc.add_paragraph(style="Pending Result")
    pending2.add_run("[PENDING ADJUDICATION] Add one evidence-linked hallucination, one distortion, one omission, one safety-focused improvement, one regression and one automated-researcher disagreement, showing exact source wording, summary wording, labels and final reason.")

    doc.add_heading("4.4 Integrated Discussion and Limitations", level=2)
    add_body(
        doc,
        "The automated findings support the literature's view that summary quality is multidimensional. The length and coverage results echo the concern that compression removes information, while the divergence between Gemini and MiniCheck reinforces SummEval's warning that automatic metrics are not interchangeable [11]. Sentence decomposition made local failures inspectable, consistent with LongEval [12], but only source-to-summary assessment exposed missing important information. The bidirectional design therefore adds information that a single faithfulness score cannot provide."
    )
    add_body(
        doc,
        "Interpretation is limited by three policies, one generation per condition and three changing API systems. Sentence rows are nested within summaries and policies rather than independent experiments. Formula-based readability does not measure understanding; the researchers are project members rather than representative users or legal experts; and Gemini's preliminary coverage denominator may reflect its own preferences. Blinding, frozen IDs, evidence requirements and matched pairs reduce some risks but do not remove subjectivity. PolicyLens demonstrates inspectability rather than legal validity or deployment safety. Final claims about hallucination, distortion, omission and automated reliability must therefore remain conditional on the pending researcher and shared-unit analyses."
    )

    words = count_words(doc)
    doc.core_properties.title = "Experimental Results and Discussion - Verifying LLM-Generated Plain-English Summaries of Privacy Policies"
    doc.core_properties.subject = "Standalone MSc AI dissertation results draft with pending human-result placeholders"
    doc.core_properties.author = "MSc AI Project Team"
    doc.core_properties.keywords = "LLM, privacy policy, matched prompt intervention, results, agreement"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"WORDS={words}")


if __name__ == "__main__":
    build()
