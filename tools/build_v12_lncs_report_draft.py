from pathlib import Path
import re

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.document import Document as _Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "outputs" / "final_report"
ASSET_DIR = OUT_DIR / "v12_draft_assets"
OUTPUT = OUT_DIR / "LNCS_V12_REPORT_DRAFT_WITH_AUTOMATED_RESULTS.docx"

INK = "111111"
MUTED = "666666"
BLUE = "244A73"
LIGHT_BLUE = "EAF1F8"
LIGHT_GRAY = "F2F2F2"
PALE_GOLD = "FFF2CC"
RED = "9C0006"
WHITE = "FFFFFF"
TABLE_WIDTH = 6912


def set_run_font(run, name="Times New Roman", size=None, bold=None, italic=None, color=None):
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    fonts.set(qn("w:ascii"), name)
    fonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=70, start=90, bottom=70, end=90):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, val in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    node = OxmlElement("w:tblHeader")
    node.set(qn("w:val"), "true")
    tr_pr.append(node)


def set_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    if tr_pr.find(qn("w:cantSplit")) is None:
        tr_pr.append(OxmlElement("w:cantSplit"))


def set_table_geometry(table, widths, indent=80):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        set_cant_split(row)
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_paragraph_left_border(paragraph, color=BLUE, size=14, space=5):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = p_bdr.find(qn("w:left"))
    if left is None:
        left = OxmlElement("w:left")
        p_bdr.append(left)
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), str(space))
    left.set(qn("w:color"), color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_run_font(run, size=8, color=MUTED)
    start = OxmlElement("w:fldChar")
    start.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([start, instr, end])


def configure_lncs_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 12, 12, 6),
        ("Heading 2", 10, 10, 4),
        ("Heading 3", 10, 8, 3),
    ):
        style = styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(INK)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = styles["Caption"]
    caption.font.name = "Times New Roman"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    caption.font.size = Pt(9)
    caption.font.bold = False
    caption.font.italic = False
    caption.font.color.rgb = RGBColor.from_string(INK)
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)
    caption.paragraph_format.keep_with_next = True

    placeholder = styles.add_style("Placeholder", 1)
    placeholder.font.name = "Times New Roman"
    placeholder._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    placeholder._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    placeholder.font.size = Pt(9.5)
    placeholder.font.bold = True
    placeholder.font.color.rgb = RGBColor.from_string(RED)
    placeholder.paragraph_format.space_before = Pt(4)
    placeholder.paragraph_format.space_after = Pt(5)
    placeholder.paragraph_format.keep_together = True

    code = styles.add_style("Prompt Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8)
    code.paragraph_format.left_indent = Cm(0.3)
    code.paragraph_format.right_indent = Cm(0.2)
    code.paragraph_format.space_before = Pt(2)
    code.paragraph_format.space_after = Pt(6)
    code.paragraph_format.line_spacing = 1.0


def add_body_paragraph(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = False
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        r1.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def add_rq(doc, label, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.45)
    p.paragraph_format.first_line_indent = Cm(-0.45)
    p.paragraph_format.keep_together = True
    r = p.add_run(f"{label}. ")
    r.bold = True
    p.add_run(text)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph(style="Caption")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run(text)
    return p


def add_table(doc, headers, rows, widths, font_size=8.2, align_numeric=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    header = table.rows[0]
    set_repeat_header(header)
    for i, text in enumerate(headers):
        shade_cell(header.cells[i], LIGHT_GRAY)
        p = header.cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(text)
        set_run_font(r, size=font_size, bold=True)
    for row_values in rows:
        row = table.add_row()
        set_cant_split(row)
        for i, value in enumerate(row_values):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_together = True
            if align_numeric and i in align_numeric:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(str(value))
            set_run_font(r, size=font_size)
    set_table_geometry(table, widths)
    return table


def add_note(doc, title, text, fill=PALE_GOLD):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.15)
    p.paragraph_format.right_indent = Cm(0.15)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_together = True
    set_paragraph_shading(p, fill)
    set_paragraph_left_border(p)
    r = p.add_run(f"{title}. ")
    r.bold = True
    p.add_run(text)
    return p


def add_prompt_block(doc, text):
    p = doc.add_paragraph(style="Prompt Block")
    set_paragraph_shading(p, "F6F6F6")
    set_paragraph_left_border(p, color="888888", size=8, space=4)
    p.paragraph_format.keep_together = False
    p.add_run(text)
    return p


def load_font(size, bold=False):
    candidates = [
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def wrap(draw, text, font, max_width):
    words = text.split()
    lines, current = [], ""
    for word in words:
        test = f"{current} {word}".strip()
        if draw.textbbox((0, 0), test, font=font)[2] <= max_width:
            current = test
        else:
            if current:
                lines.append(current)
            current = word
    if current:
        lines.append(current)
    return lines


def draw_box(draw, xy, title, body, fill, outline=BLUE):
    x1, y1, x2, y2 = xy
    draw.rounded_rectangle(xy, radius=18, fill=f"#{fill}", outline=f"#{outline}", width=4)
    title_font = load_font(30, bold=True)
    body_font = load_font(24)
    title_lines = wrap(draw, title, title_font, x2 - x1 - 40)
    body_lines = wrap(draw, body, body_font, x2 - x1 - 40)
    y = y1 + 22
    for line in title_lines:
        box = draw.textbbox((0, 0), line, font=title_font)
        draw.text(((x1 + x2 - (box[2] - box[0])) / 2, y), line, fill=f"#{INK}", font=title_font)
        y += 36
    y += 8
    for line in body_lines:
        box = draw.textbbox((0, 0), line, font=body_font)
        draw.text(((x1 + x2 - (box[2] - box[0])) / 2, y), line, fill=f"#{INK}", font=body_font)
        y += 30


def arrow(draw, start, end):
    draw.line([start, end], fill=f"#{BLUE}", width=8)
    x, y = end
    draw.polygon([(x, y), (x - 16, y - 25), (x + 16, y - 25)], fill=f"#{BLUE}")


def make_vertical_diagram(path, title, boxes):
    width = 1500
    box_h = 170
    gap = 70
    top = 130
    height = top + len(boxes) * box_h + (len(boxes) - 1) * gap + 80
    img = Image.new("RGB", (width, height), f"#{WHITE}")
    draw = ImageDraw.Draw(img)
    title_font = load_font(46, bold=True)
    draw.text((75, 35), title, fill=f"#{BLUE}", font=title_font)
    draw.line((75, 100, width - 75, 100), fill=f"#{BLUE}", width=5)
    left, right = 190, width - 190
    y = top
    for i, (box_title, body, fill) in enumerate(boxes):
        draw_box(draw, (left, y, right, y + box_h), box_title, body, fill)
        if i < len(boxes) - 1:
            arrow(draw, (width // 2, y + box_h), (width // 2, y + box_h + gap - 10))
        y += box_h + gap
    img.save(path, dpi=(180, 180))


def build_diagrams():
    ASSET_DIR.mkdir(parents=True, exist_ok=True)
    make_vertical_diagram(
        ASSET_DIR / "figure_B1_v12_pipeline.png",
        "Experimental and evaluation pipeline",
        [
            ("Frozen corpus", "Three complete policies: short, medium and long/complex", "EAF1F8"),
            ("Matched generation", "3 models x 3 strategies x 2 prompt sets = 54 summaries", "EAF1F8"),
            ("Automated evaluation", "Readability, compression, Gemini verification and MiniCheck", "F2F2F2"),
            ("Researcher evaluation", "Three research-team members; anonymised independent assessment", "F2F2F2"),
            ("Analysis and prototype", "Matched intervention effects, agreement and evidence inspection", "D9E7F5"),
        ],
    )
    make_vertical_diagram(
        ASSET_DIR / "figure_B2_v12_faithfulness.png",
        "Summary-to-source verification",
        [
            ("Frozen summary sentence", "1,393 sentences across 54 blinded summaries", "EAF1F8"),
            ("Complete source policy", "Evidence is checked against the full frozen policy", "EAF1F8"),
            ("Independent judgements", "Gemini: four labels; MiniCheck: binary; A1-A3: four labels", "F2F2F2"),
            ("Evidence and reason", "Exact source wording, explanation and confidence where applicable", "F2F2F2"),
            ("Agreement and adjudication", "Compare identical IDs; retain distributions and resolve disagreements", "D9E7F5"),
        ],
    )
    make_vertical_diagram(
        ASSET_DIR / "figure_B3_v12_coverage.png",
        "Source-to-summary coverage verification",
        [
            ("Read source first", "Identify important information actually present in each policy", "EAF1F8"),
            ("Independent proposals", "A1, A2 and A3 work without summaries or automated outputs", "EAF1F8"),
            ("Adjudicate and freeze", "Merge duplicates, preserve qualifications and create one shared unit set", "F2F2F2"),
            ("Check every summary", "Covered, Partially covered or Not covered, with evidence and reason", "F2F2F2"),
            ("Human-Gemini agreement", "Score the identical unit-summary pairs and analyse omission", "D9E7F5"),
        ],
    )


def iter_block_items(parent):
    if isinstance(parent, _Document):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("Unsupported parent")
    for child in parent_elm.iterchildren():
        if child.tag == qn("w:p"):
            yield Paragraph(child, parent)
        elif child.tag == qn("w:tbl"):
            yield Table(child, parent)


def word_count_between(doc, start_heading, end_heading):
    active = False
    text = []
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            value = block.text.strip()
            if value == start_heading:
                active = True
            if value == end_heading:
                active = False
                break
            if active:
                text.append(value)
        elif active:
            for row in block.rows:
                text.extend(cell.text for cell in row.cells)
    return len(re.findall(r"\b[\w'-]+\b", " ".join(text)))


def build_report():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    build_diagrams()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(5.2)
    section.bottom_margin = Cm(5.2)
    section.left_margin = Cm(4.4)
    section.right_margin = Cm(4.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)
    configure_lncs_styles(doc)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(10)
    title.paragraph_format.keep_with_next = True
    r = title.add_run("Verifying LLM-Generated Plain-English Summaries of Privacy Policies")
    set_run_font(r, size=16, bold=True)

    authors = doc.add_paragraph()
    authors.alignment = WD_ALIGN_PARAGRAPH.CENTER
    authors.paragraph_format.space_after = Pt(3)
    authors.paragraph_format.keep_with_next = True
    ar = authors.add_run("[Author 1 name and UWE ID]  |  [Author 2 name and UWE ID]  |  [Author 3 name and UWE ID]")
    set_run_font(ar, size=10)

    affiliation = doc.add_paragraph()
    affiliation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    affiliation.paragraph_format.space_after = Pt(5)
    affiliation.paragraph_format.keep_with_next = True
    afr = affiliation.add_run("University of the West of England, Bristol, UK")
    set_run_font(afr, size=9, italic=True)

    status = doc.add_paragraph(style="Placeholder")
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    status.add_run("WORKING DRAFT: update author details, human results and final GitHub URL before submission.")

    abstract_text = (
        "Large language models can make privacy policies shorter and more accessible, but fluent output may add unsupported information, alter source meaning or omit information important to a reader. This controlled case study generated 54 plain-English summaries from three policies using three model families, three prompting strategies and matched basic and safety-focused prompt sets. Evaluation combines readability and compression with bidirectional verification: summary sentences are checked for hallucination and distortion, while source information is checked for omission. Gemini and MiniCheck completed automated assessment of all 1,393 summary sentences. Safety-focused prompting most clearly improved Gemini-assessed information coverage, but produced longer, generally harder-to-read summaries and did not improve MiniCheck support consistently. Three blinded student researchers are evaluating the same material; therefore final human agreement and coverage conclusions remain pending. The study contributes an evidence-linked evaluation workflow and an interactive verification prototype."
    )
    abstract = doc.add_paragraph()
    abstract.paragraph_format.space_after = Pt(5)
    abstract.paragraph_format.keep_together = True
    rr = abstract.add_run("Abstract. ")
    set_run_font(rr, size=9, bold=True)
    rr2 = abstract.add_run(abstract_text)
    set_run_font(rr2, size=9)

    keywords = doc.add_paragraph()
    keywords.paragraph_format.space_after = Pt(10)
    kr = keywords.add_run("Keywords: ")
    set_run_font(kr, size=9, bold=True)
    kr2 = keywords.add_run("large language models; privacy-policy summarisation; plain English; source alignment; human evaluation")
    set_run_font(kr2, size=9)

    doc.add_heading("1 Introduction", level=1)
    doc.add_heading("1.1 Context and Rationale", level=2)
    add_body_paragraph(doc, "Privacy policies explain how organisations collect, use, share, retain and protect personal information. They influence whether people create accounts, provide data or exercise privacy rights, yet the documents are often long and legally qualified. McDonald and Cranor quantified the collective reading burden of online policies [1], and a later million-document analysis found that policies had grown longer and more difficult to read over time [2]. Formal publication therefore does not guarantee practical access for an ordinary reader.")
    add_body_paragraph(doc, "Regulatory guidance requires privacy information to be concise, transparent, intelligible and written in clear language [3]. Large language models (LLMs) offer a practical way to condense and rephrase full policies for non-specialists. However, a fluent summary can still be unreliable. It may add an unsupported statement, alter a material qualification such as 'may', 'only' or 'unless', or omit information a reader would need. Readability and reliability must consequently be evaluated together rather than assuming that simpler wording preserves meaning automatically.")

    doc.add_heading("1.2 Aim, Scope and Contribution", level=2)
    add_body_paragraph(doc, "This project investigates whether contemporary general-purpose LLMs can produce useful plain-English summaries of complete privacy policies and how those summaries can be verified. It narrows the original legal-document brief to a controlled privacy-policy case study: three policies representing short, medium and long/complex documents are processed by three model families, three prompting strategies and two matched prompt sets. The resulting 54 summaries support detailed evidence-linked evaluation, but not population-wide claims about all policies, LLMs or readers.")
    add_body_paragraph(doc, "The contribution is an inspectable verification framework rather than a new language model. It combines formula-based accessibility indicators with two directions of source alignment. Summary-to-source assessment checks generated sentences for hallucination and distortion; source-to-summary assessment checks whether independently identified important information has been omitted. Automated judgements are compared with three blinded student researchers, and a Streamlit prototype exposes the labels, source evidence and reasons behind individual decisions.")

    doc.add_heading("1.3 Research Questions", level=2)
    add_rq(doc, "RQ1", "To what extent can the selected LLMs generate privacy-policy summaries that are readable, concise and aligned with the source?")
    add_rq(doc, "RQ2", "How effectively does safety-focused prompting reduce hallucination, distortion and omission while maintaining readability and conciseness?")
    add_rq(doc, "RQ3", "To what extent do automated evaluator judgements align with researcher judgements when identifying hallucination, distortion and omission?")

    doc.add_heading("1.4 Report Structure", level=2)
    add_body_paragraph(doc, "Section 2 reviews privacy-policy accessibility, legal summarisation and evaluation research. Section 3 describes the corpus, matched prompt intervention, automated and researcher assessment, and prototype. Section 4 reports completed automated results and reserves clearly marked locations for the pending human analyses. Section 5 gives provisional conclusions and next steps.")

    doc.add_heading("2 Relevant Work", level=1)
    doc.add_heading("2.1 Privacy Policies and Plain Language", level=2)
    add_body_paragraph(doc, "Privacy policies are intended to make organisational data practices transparent, but length, specialised vocabulary and layered legal qualifications can obstruct meaningful access [1, 2]. The Information Commissioner's Office recommends short sentences, common language, precision and information that is easy to locate; it also recommends testing privacy information with its intended users [3]. This distinction is important: formula-based readability reflects textual characteristics, not whether a reader actually understands a data practice or whether simplification remains legally adequate.")
    add_body_paragraph(doc, "Earlier privacy NLP mainly transformed policies into structured representations. OPP-115 supplied expert annotations of privacy practices across 115 policies [4]. Polisis used privacy-specific modelling and hierarchical classification to support scalable queries [5], while PrivacyQA represented access as question answering and reported a sizeable gap between a strong neural baseline and human performance [6]. These systems improve navigation, but classification and retrieval differ from open-ended plain-English generation.")

    doc.add_heading("2.2 Privacy and Legal Summarisation", level=2)
    add_body_paragraph(doc, "A review of consumer-privacy NLP describes extractive methods that select relevant or risky sentences and emphasises that choosing what users should see is itself difficult [7]. EROS later used entity-driven controlled abstractive summarisation to improve the readability and information content of policy summaries [8]. These studies show that privacy-policy summarisation is not new; the unresolved issue is whether flexible generation preserves both the meaning and breadth of the source.")
    add_body_paragraph(doc, "Evidence from broader legal summarisation reinforces this concern. LexSumm reports abstraction and faithfulness errors in zero-shot LLM summaries across eight legal datasets [9]. CaseSumm found disagreements between automatic metrics and legal-expert assessment: summaries favoured by some automatic measures could still contain hallucinations or misrepresent case facts [10]. Although court opinions differ from consumer notices, both domains show why fluent legal summaries require source-linked validation.")

    doc.add_heading("2.3 Reliability Failures and Prompting", level=2)
    add_body_paragraph(doc, "This study distinguishes three failures. Hallucination is material content for which adequate source evidence cannot be found. Distortion is content with a source basis whose certainty, scope, actor, condition, exception or relationship has materially changed. Omission is important information that is present in the source but not adequately retained. Faithfulness and coverage are complementary: a summary can contain only supported sentences yet remain misleading because it excludes central practices.")
    add_body_paragraph(doc, "Prompting may influence these outcomes without retraining a model. Minimal direct instructions, an audience-oriented role and a structured internal planning instruction can change selection and wording. A second, safety-focused prompt can explicitly prohibit assumptions and require preservation of qualifications. However, added safeguards may increase length, copy source complexity or encourage elaboration. Prompt design is therefore treated as a controlled intervention, not assumed to improve every quality dimension.")

    doc.add_heading("2.4 Evaluation and Research Gap", level=2)
    add_body_paragraph(doc, "SummEval demonstrated that automatic summary metrics capture different dimensions and correspond imperfectly with human judgement [11]. For long-form faithfulness, LongEval found that finer-grained judgement reduced annotator variance [12], supporting sentence-level rather than whole-summary labels. G-Eval reported improved human correspondence for an LLM-based judge but also raised concerns about evaluator bias [13]. MiniCheck provides a smaller document-grounded fact checker [14], yet meta-evaluation is still dominated by news datasets [15], and STORYSUMM shows that both human protocols and automatic metrics can miss difficult inconsistencies [16].")
    add_body_paragraph(doc, "The gap is consequently not whether policies can be summarised. It is the limited controlled evidence on whether general-purpose LLMs can balance readability and conciseness with accurate and sufficiently broad preservation of full-policy information, and whether automated verification is credible in this domain. This study addresses that gap through matched prompt intervention, bidirectional evidence checking and automated-researcher agreement analysis.")

    doc.add_heading("3 Methods and Materials", level=1)
    doc.add_heading("3.1 Study Design and Corpus", level=2)
    add_body_paragraph(doc, "A controlled within-document design crossed policy, model family, prompting strategy and prompt set. Three publicly available English-language policies were purposively selected from the frozen pilot corpus to represent short, medium and long/complex cases. Original webpages, cleaned text, metadata and cryptographic hashes were retained separately. Cleaning removed duplicated webpage furniture and malformed spacing while preserving substantive wording, headings, conditions and exceptions.")
    add_caption(doc, "Table 1. Frozen privacy-policy corpus.")
    add_table(
        doc,
        ["ID", "Organisation", "Role", "Words", "Retrieved"],
        [
            ("PILOT01", "Mozilla", "Short", "857", "12 Aug 2026"),
            ("PILOT02", "DuckDuckGo", "Medium", "2,077", "12 Aug 2026"),
            ("PILOT03", "Automattic", "Long/complex", "8,285", "12 Aug 2026"),
        ],
        [1050, 1750, 1450, 900, 1762],
        font_size=8.1,
        align_numeric={3, 4},
    )
    add_body_paragraph(doc, "Each policy was summarised under 3 models x 3 strategies x 2 prompt sets, producing 18 outputs per policy and 54 outputs overall. Every safety-focused output was paired with the basic output generated from the same policy, model and strategy, giving 27 matched intervention pairs. Figure B1 summarises the complete workflow.")

    doc.add_heading("3.2 Models, Prompts and Generation", level=2)
    add_body_paragraph(doc, "The frozen systems were GPT-5.6 Luna (OpenAI; gpt-5.6-luna), Llama 3.3 70B Instruct Turbo (Together AI; meta-llama/Llama-3.3-70B-Instruct-Turbo) and Mistral Large 3 (Mistral API; mistral-large-2512). They provide distinct model families and provider ecosystems available through comparable APIs. The selection is not claimed to be a complete open-versus-closed benchmark because training data, architecture, scale and deployment also differ.")
    add_caption(doc, "Table 2. Prompting conditions in the matched intervention.")
    add_table(
        doc,
        ["Factor", "Condition", "Operational distinction"],
        [
            ("Strategy", "Direct", "Task and audience stated with minimal guidance"),
            ("Strategy", "Role-guided", "Adds a plain-language communication role"),
            ("Strategy", "Structured", "Adds private planning and logical organisation; final summary only"),
            ("Prompt set", "Basic", "Ordinary summarisation instruction"),
            ("Prompt set", "Safety-focused", "Adds source-only, qualification-preservation and information-retention safeguards"),
        ],
        [1150, 1500, 4262],
        font_size=8.1,
    )
    add_body_paragraph(doc, "All prompts specified a general adult reader without specialist legal or privacy knowledge and required continuous prose, complete sentences and short paragraphs. The safety-focused set preserved the same task and format while adding safeguards against unsupported assumptions, changed certainty, lost exceptions and missing important information. Full frozen templates are reproduced in Appendix C.")
    add_body_paragraph(doc, "One output was generated per condition. Llama and Mistral used temperature 0.0; the GPT endpoint did not support the temperature parameter, so it was omitted and logged as a provider constraint. The maximum output allowance was 6,000 tokens. API logs recorded model identifiers, timestamps, token usage, status and errors. Valid but poor outputs were retained rather than regenerated; one Mistral structured output used bullets despite the prose constraint and was preserved as observed behaviour.")

    doc.add_heading("3.3 Accessibility and Conciseness", level=2)
    add_body_paragraph(doc, "Readability was estimated at summary level using Flesch Reading Ease, Flesch-Kincaid Grade and SMOG Grade. Conciseness was represented by word count and word-compression ratio, defined as summary words divided by source-policy words. Higher Reading Ease and lower grade estimates conventionally indicate simpler textual structure, while a lower compression ratio indicates a shorter summary relative to its source. These measures do not demonstrate comprehension, legal accuracy or usability; formatting and sentence segmentation may also affect them.")

    doc.add_heading("3.4 Bidirectional Source Alignment", level=2)
    add_body_paragraph(doc, "For summary-to-source verification, summaries were deterministically divided into visible sentences. Gemini 3.5 Flash Lite received each sentence, the complete frozen policy and summary context, then assigned Supported, Partially supported, Contradicted or Unsupported, together with a failure type, exact source evidence and a reason. Partially supported normally indicates distortion; Unsupported normally indicates hallucination. MiniCheck with FLAN-T5-Large independently returned a binary support prediction and probability for the identical sentence IDs. MiniCheck was not converted into four categories and was not used to measure omission. Figure B2 shows this direction of verification.")
    add_body_paragraph(doc, "Coverage used the reverse direction. Gemini first proposed policy-specific information important to a non-specialist reader, without applying a fixed privacy taxonomy. Evidence validation retained 20 units for PILOT01, 19 for PILOT02 and 13 for PILOT03. Each unit was then compared with the 18 summaries associated with its policy and labelled Covered, Partially covered or Not covered. Strict coverage equals Covered divided by all units; weighted coverage equals (Covered + 0.5 x Partially covered) divided by all units. The 0.5 weight is a transparent convention, not a natural ground truth.")

    doc.add_heading("3.5 Researcher Evaluation and Analysis", level=2)
    add_body_paragraph(doc, "Three student researchers receive identical blinded packages containing the three complete policies, 54 summaries and 1,393 frozen sentence rows. They work independently and cannot view model identities, automated results or one another's decisions. For every sentence they record a label, matching failure type, exact evidence, reason and confidence. They also identify important source information before reading the summaries. After structural validation, proposals will be adjudicated into one shared policy-specific coverage denominator; all researchers and Gemini will then score the same unit-summary comparisons (Fig. B3). An absent topic cannot be counted as omitted unless the policy contains it and it enters the adjudicated unit set.")
    add_body_paragraph(doc, "Outcomes are described by condition. RQ2 uses safety-focused minus basic differences within each matched policy-model-strategy pair. Because only three policies were sampled, analysis emphasises magnitude, direction and consistency rather than population-wide significance. Researcher reliability will be reported through exact agreement, Cohen's kappa, category distributions and confusion matrices. Written evidence and reasons support qualitative analysis of recurring failure and disagreement patterns.")

    doc.add_heading("3.6 Verification Prototype", level=2)
    add_body_paragraph(doc, "A Streamlit prototype named PolicyLens demonstrates the proposed verification workflow. Users select a policy, model, strategy and prompt set; click a summary sentence to view the Gemini label, failure type, exact policy evidence, explanation and MiniCheck result; inspect source-first coverage units; and explore evaluator disagreements. Green, amber, red and grey consistently represent supported/covered, partial/distorted, unsupported/omitted and pending states. The prototype is an inspectable research interface, not a legal-advice service.")

    doc.add_heading("4 Experimental Results and Discussion", level=1)
    doc.add_heading("4.1 Completion and Automated Evaluator Distributions", level=2)
    add_body_paragraph(doc, "All 54 planned summaries were generated successfully and formed 27 complete matched pairs. Gemini and MiniCheck each completed all 1,393 frozen sentence assessments. Gemini labelled 1,345 sentences Supported, 34 Partially supported, 13 Unsupported and one Contradicted; its failure taxonomy contained 34 distortions, 13 hallucinations and one combined hallucination-and-distortion case. MiniCheck predicted 1,255 sentences supported and 138 unsupported. These distributions differ substantially and are not interchangeable estimates of factual accuracy.")
    add_body_paragraph(doc, "Gemini also completed 936 preliminary source-unit-summary comparisons: 550 Covered, 192 Partially covered and 194 Not covered. These figures demonstrate why sentence faithfulness cannot establish completeness. They remain preliminary because the units were proposed by Gemini; final omission and human-Gemini agreement will use the shared human-adjudicated denominator.")

    doc.add_heading("4.2 RQ1: Accessibility, Conciseness and Automated Source Alignment", level=2)
    add_caption(doc, "Table 3. Mean automated outcomes by prompt set (27 summaries per row).")
    add_table(
        doc,
        ["Prompt set", "Words", "FRE", "FK", "SMOG", "Comp.", "Gem. sup.", "Mini. sup.", "Prelim. cov."],
        [
            ("Basic", "480.9", "38.11", "13.31", "15.48", "0.223", "94.38%", "91.84%", "48.90%"),
            ("Safety", "662.4", "35.90", "14.12", "16.08", "0.304", "96.75%", "90.74%", "64.19%"),
        ],
        [1000, 680, 650, 650, 650, 720, 820, 820, 922],
        font_size=7.3,
        align_numeric={1, 2, 3, 4, 5, 6, 7, 8},
    )
    add_body_paragraph(doc, "The summaries were much shorter than their sources, retaining on average 22.3% of source words under basic prompting and 30.4% under safety-focused prompting. However, mean Reading Ease remained below 40 and grade estimates remained above 13, so the outputs were not demonstrably easy for a general adult reader. These formulae cannot determine actual understanding, but they do not support a strong claim that the summaries achieved universally simple language.")
    add_body_paragraph(doc, "Descriptively across both prompt sets, GPT produced the longest summaries (835.4 words) and the highest preliminary strict coverage (84.8%); Llama produced the shortest (352.2 words) and lowest preliminary coverage (34.4%); Mistral produced the highest Reading Ease (43.85) and lowest Flesch-Kincaid Grade (12.57). These patterns suggest a length-coverage trade-off and possible family differences, but each mean is based on only three policies and should not be generalised beyond the controlled cases.")

    doc.add_heading("4.3 RQ2: Effect of Safety-Focused Prompting", level=2)
    add_caption(doc, "Table 4. Safety-focused minus basic differences across 27 matched pairs.")
    add_table(
        doc,
        ["Outcome", "Mean delta", "Median", "+ / 0 / - pairs"],
        [
            ("Words", "+181.56", "+136.00", "24 / 0 / 3"),
            ("Flesch Reading Ease", "-2.21", "-3.18", "8 / 0 / 19"),
            ("Flesch-Kincaid Grade", "+0.81", "+0.81", "20 / 0 / 7"),
            ("Gemini strict support", "+2.37 pp", "+1.21 pp", "16 / 4 / 7"),
            ("MiniCheck support", "-1.09 pp", "-0.55 pp", "10 / 3 / 14"),
            ("Preliminary strict coverage", "+15.29 pp", "+15.79 pp", "22 / 4 / 1"),
            ("Preliminary weighted coverage", "+13.02 pp", "+15.00 pp", "26 / 0 / 1"),
        ],
        [2650, 1450, 1250, 1562],
        font_size=8.0,
        align_numeric={1, 2, 3},
    )
    add_body_paragraph(doc, "The intervention did not improve every dimension. Safety-focused outputs were longer in 24 of 27 pairs and had lower Reading Ease in 19. Gemini strict support increased by 2.37 percentage points on average, but MiniCheck support decreased by 1.09 points and declined in 14 pairs. The apparent faithfulness benefit therefore depends on the evaluator, reinforcing the need for RQ3 rather than selecting the more favourable automated score.")
    add_body_paragraph(doc, "The clearest completed effect was preliminary information retention. Strict coverage increased by 15.29 percentage points and improved in 22 pairs; weighted coverage improved in 26 pairs. The safeguards appear to have encouraged broader preservation of source information, but at the cost of conciseness and formula-based readability. This result is consistent with the theoretical tension identified in Section 2: asking a model to preserve qualifications and important content can reduce omission while reproducing more of the source's complexity.")

    doc.add_heading("4.4 RQ3: Automated-Researcher Agreement", level=2)
    pending = doc.add_paragraph(style="Placeholder")
    pending.add_run("[PENDING HUMAN RESULTS] Insert validated A1/A2/A3 label distributions, pairwise exact agreement and Cohen's kappa, adjudicated sentence labels, Gemini-human four-class and binary comparisons, MiniCheck-human binary comparison, and shared-unit coverage agreement.")
    add_body_paragraph(doc, "The completed Gemini-MiniCheck binary comparison can be displayed in the prototype as a diagnostic, but agreement between two automated systems is not validation. RQ3 can only be answered after comparison with the independent researcher decisions on identical sentence and coverage IDs. Confusion matrices and prevalence will accompany agreement coefficients because a high overall value may be dominated by the Supported class.")

    doc.add_heading("4.5 Qualitative Failure Analysis and Prototype", level=2)
    pending2 = doc.add_paragraph(style="Placeholder")
    pending2.add_run("[PENDING ADJUDICATION] Insert one evidence-linked hallucination, one altered qualification or scope, one omission, one safety-focused improvement, one regression and one automated-evaluator disagreement. Include exact source wording, summary wording, labels and final reason.")
    add_body_paragraph(doc, "The prototype already supports this analysis by connecting each label to source evidence and surrounding context. Its value is transparency: a reader can inspect why a sentence was considered supported or distorted rather than accepting an unexplained aggregate score. A final report screenshot should be inserted after human labels are loaded so the demonstration represents the completed system.")

    doc.add_heading("4.6 Discussion and Limitations", level=2)
    add_body_paragraph(doc, "The automated findings do not support a simple claim that one prompt set solved legal-summary reliability. Safety-focused prompting substantially improved Gemini-assessed coverage, slightly improved Gemini support, worsened formula-based accessibility and did not improve MiniCheck consistently. This multidimensional trade-off echoes SummEval's argument for comprehensive evaluation [11] and CaseSumm's warning that automatic and expert assessments may diverge in legal summarisation [10].")
    add_body_paragraph(doc, "The study is limited to three purposively selected policies, one generation per condition and three model families accessed through changing APIs. Sentence rows are nested within summaries and policies and cannot be treated as independent experiments. Readability formulae do not measure comprehension; the three researchers are project members rather than representative end users or legal experts; and Gemini's preliminary coverage units may reflect its own preferences. Blinding, evidence requirements and shared-unit adjudication reduce some risks but do not remove subjectivity. The prototype demonstrates inspectability rather than deployment safety or legal validity.")

    doc.add_heading("5 Conclusions", level=1)
    add_body_paragraph(doc, "This controlled case study reframes LLM comparison as one component of a wider verification problem. The 54 summaries were substantially shorter than the policies, but formula-based results did not establish uniformly accessible plain English. Automated sentence assessment was generally favourable, although Gemini and MiniCheck differed in both distribution and intervention direction. Safety-focused prompting most clearly improved preliminary information coverage, while producing longer and generally harder-to-read summaries. It therefore created a trade-off rather than a universal quality improvement.")
    add_body_paragraph(doc, "A final answer to RQ3, and a confirmed conclusion about hallucination, distortion and omission, must wait for the blinded researcher and shared-unit coverage analyses. Subject to that validation, the practical conclusion is that fluent legal summaries should not be trusted on presentation alone: verification must examine what the summary says, whether material meaning changed and what important source information is absent.")
    add_body_paragraph(doc, "Future work should expand the policy corpus, repeat stochastic generations, involve independent legal specialists and intended users, and measure task-based comprehension. The evidence-linked prototype could also be extended from displaying failures to warning about, revising or regenerating unsupported and incomplete passages.")

    doc.add_heading("References", level=1)
    references = [
        "[1] McDonald, A.M., Cranor, L.F.: The cost of reading privacy policies. I/S: A Journal of Law and Policy for the Information Society 4(3), 543-568 (2008).",
        "[2] Amos, R., Acar, G., Lucherini, E., Kshirsagar, M., Narayanan, A., Mayer, J.: Privacy policies over time: Curation and analysis of a million-document dataset. In: The Web Conference 2021, pp. 2165-2176 (2021). https://doi.org/10.1145/3442381.3450048",
        "[3] Information Commissioner's Office: How should we draft our privacy information? https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/individual-rights/the-right-to-be-informed/how-should-we-draft-our-privacy-information/, last accessed 2026/08/23.",
        "[4] Wilson, S., et al.: The creation and analysis of a website privacy policy corpus. In: ACL 2016, pp. 1330-1340 (2016). https://doi.org/10.18653/v1/P16-1126",
        "[5] Harkous, H., Fawaz, K., Lebret, R., Schaub, F., Shin, K.G., Aberer, K.: Polisis: Automated analysis and presentation of privacy policies using deep learning. In: USENIX Security 2018, pp. 531-548 (2018).",
        "[6] Ravichander, A., Black, A.W., Wilson, S., Norton, T., Sadeh, N.: Question answering for privacy policies: Combining computational and legal perspectives. In: EMNLP-IJCNLP 2019, pp. 4947-4958 (2019). https://doi.org/10.18653/v1/D19-1500",
        "[7] Ravichander, A., Black, A.W., Norton, T., Wilson, S., Sadeh, N.: Breaking down walls of text: How can NLP benefit consumer privacy? In: ACL-IJCNLP 2021, pp. 4125-4140 (2021). https://doi.org/10.18653/v1/2021.acl-long.319",
        "[8] Singh, J., Fazili, S., Jain, R., Akhtar, M.S.: EROS: Entity-driven controlled policy document summarization. In: LREC-COLING 2024, pp. 6236-6246 (2024).",
        "[9] T.Y.S.S, S., Weiss, C., Grabmair, M.: LexSumm and LexT5: Benchmarking and modeling legal summarization tasks in English. In: NLLP 2024, pp. 381-403 (2024). https://doi.org/10.18653/v1/2024.nllp-1.35",
        "[10] Heddaya, M., MacMillan, K., Mei, H., Tan, C., Malani, A.: CaseSumm: A large-scale dataset for long-context summarization from U.S. Supreme Court opinions. In: Findings of NAACL 2025, pp. 1917-1942 (2025). https://doi.org/10.18653/v1/2025.findings-naacl.102",
        "[11] Fabbri, A.R., Kryscinski, W., McCann, B., Xiong, C., Socher, R., Radev, D.: SummEval: Re-evaluating summarization evaluation. TACL 9, 391-409 (2021). https://doi.org/10.1162/tacl_a_00373",
        "[12] Krishna, K., et al.: LongEval: Guidelines for human evaluation of faithfulness in long-form summarization. In: EACL 2023, pp. 1650-1669 (2023). https://doi.org/10.18653/v1/2023.eacl-main.121",
        "[13] Liu, Y., Iter, D., Xu, Y., Wang, S., Xu, R., Zhu, C.: G-Eval: NLG evaluation using GPT-4 with better human alignment. In: EMNLP 2023, pp. 2511-2522 (2023). https://doi.org/10.18653/v1/2023.emnlp-main.153",
        "[14] Tang, L., Laban, P., Durrett, G.: MiniCheck: Efficient fact-checking of LLMs on grounding documents. In: EMNLP 2024, pp. 8818-8847 (2024). https://doi.org/10.18653/v1/2024.emnlp-main.499",
        "[15] Dai, X., Karimi, S., Fang, B.: A critical look at meta-evaluating summarisation evaluation metrics. In: Findings of EMNLP 2024, pp. 14795-14808 (2024). https://doi.org/10.18653/v1/2024.findings-emnlp.869",
        "[16] Subbiah, M., Ladhak, F., Mishra, A., Adams, G.T., Chilton, L., McKeown, K.: STORYSUMM: Evaluating faithfulness in story summarization. In: EMNLP 2024, pp. 9988-10005 (2024). https://doi.org/10.18653/v1/2024.emnlp-main.557",
    ]
    for ref in references:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.first_line_indent = Cm(-0.45)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(ref)

    doc.add_heading("Appendix A: Code Repository", level=1)
    p = doc.add_paragraph(style="Placeholder")
    p.add_run("GitHub repository: [INSERT FINAL PUBLIC URL, RELEASE/TAG AND COMMIT HASH].")

    doc.add_heading("Appendix B: Modelling Diagrams", level=1)
    for image_name, caption in (
        ("figure_B1_v12_pipeline.png", "Fig. B1. Frozen v1.2 experimental and evaluation pipeline."),
        ("figure_B2_v12_faithfulness.png", "Fig. B2. Summary-to-source verification for hallucination and distortion."),
        ("figure_B3_v12_coverage.png", "Fig. B3. Source-to-summary coverage workflow for omission."),
    ):
        pic = doc.add_paragraph()
        pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pic.paragraph_format.keep_with_next = True
        shape = pic.add_run().add_picture(str(ASSET_DIR / image_name), width=Cm(12.0))
        doc_pr = shape._inline.docPr
        doc_pr.set("title", caption.split(". ", 1)[0])
        doc_pr.set("descr", caption.split(". ", 1)[1].rstrip("."))
        add_caption(doc, caption)

    doc.add_heading("Appendix C: Frozen Prompt Templates", level=1)
    add_note(doc, "Reporting note", "The full policy text replaced {{POLICY_TEXT}} at generation time. No target word range was imposed in v1.2.", LIGHT_BLUE)
    prompts = [
        ("C.1 Basic direct", "Summarise the privacy policy below in clear, concise plain English for a general adult reader with no specialist legal or privacy knowledge.\n\nWrite the summary as continuous prose using complete sentences and short paragraphs. Do not use headings, bullet points, tables or numbered lists.\n\nReturn only the summary.\n\nPRIVACY POLICY:\n\n{{POLICY_TEXT}}"),
        ("C.2 Basic role-guided", "You are a plain-language communication specialist who explains complex information to non-specialist readers.\n\nSummarise the privacy policy below in clear, concise plain English for a general adult reader with no specialist legal or privacy knowledge.\n\nWrite the summary as continuous prose using complete sentences and short paragraphs. Do not use headings, bullet points, tables or numbered lists.\n\nReturn only the summary.\n\nPRIVACY POLICY:\n\n{{POLICY_TEXT}}"),
        ("C.3 Basic structured", "Summarise the privacy policy below in clear, concise plain English for a general adult reader with no specialist legal or privacy knowledge.\n\nBefore writing, internally identify the policy information most important to the intended reader, organise it logically, and plan the summary. Do not reveal this process.\n\nWrite the summary as continuous prose using complete sentences and short paragraphs. Do not use headings, bullet points, tables or numbered lists.\n\nReturn only the summary.\n\nPRIVACY POLICY:\n\n{{POLICY_TEXT}}"),
        ("C.4 Safety-focused direct", "Summarise the privacy policy below in clear, concise plain English for a general adult reader with no specialist legal or privacy knowledge.\n\nRetain the important information the reader needs to understand how the organisation handles personal information, what consequences this may have, and what rights or choices are available.\n\nUse only information supported by the policy. Preserve the original meaning, including important conditions, exceptions, limitations, quantities, time periods and distinctions such as 'may,' 'will' and 'only.' Do not add assumptions or present uncertain information as certain.\n\nWrite the summary as continuous prose using complete sentences and short paragraphs. Do not use headings, bullet points, tables or numbered lists.\n\nReturn only the summary.\n\nPRIVACY POLICY:\n\n{{POLICY_TEXT}}"),
        ("C.5 Safety-focused role-guided", "You are a plain-language communication specialist who explains complex information to non-specialist readers.\n\nSummarise the privacy policy below in clear, concise plain English for a general adult reader with no specialist legal or privacy knowledge.\n\nRetain the important information the reader needs to understand how the organisation handles personal information, what consequences this may have, and what rights or choices are available.\n\nUse only information supported by the policy. Preserve the original meaning, including important conditions, exceptions, limitations, quantities, time periods and distinctions such as 'may,' 'will' and 'only.' Do not add assumptions or present uncertain information as certain.\n\nWrite the summary as continuous prose using complete sentences and short paragraphs. Do not use headings, bullet points, tables or numbered lists.\n\nReturn only the summary.\n\nPRIVACY POLICY:\n\n{{POLICY_TEXT}}"),
        ("C.6 Safety-focused structured", "Summarise the privacy policy below in clear, concise plain English for a general adult reader with no specialist legal or privacy knowledge.\n\nBefore writing, internally identify the policy information most important to the intended reader, organise it logically, and plan the summary. Do not reveal this process.\n\nRetain the important information the reader needs to understand how the organisation handles personal information, what consequences this may have, and what rights or choices are available.\n\nUse only information supported by the policy. Preserve the original meaning, including important conditions, exceptions, limitations, quantities, time periods and distinctions such as 'may,' 'will' and 'only.' Do not add assumptions or present uncertain information as certain.\n\nWrite the summary as continuous prose using complete sentences and short paragraphs. Do not use headings, bullet points, tables or numbered lists.\n\nReturn only the summary.\n\nPRIVACY POLICY:\n\n{{POLICY_TEXT}}"),
    ]
    for heading, prompt in prompts:
        doc.add_heading(heading, level=2)
        add_prompt_block(doc, prompt)

    body_words = word_count_between(doc, "1 Introduction", "References")
    abstract_words = len(re.findall(r"\b[\w'-]+\b", abstract_text))
    doc.core_properties.title = "Verifying LLM-Generated Plain-English Summaries of Privacy Policies"
    doc.core_properties.subject = "MSc AI dissertation working draft - v1.2 study"
    doc.core_properties.author = "MSc AI Project Team"
    doc.core_properties.keywords = "LLM, privacy policy, plain English, faithfulness, coverage"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"BODY_WORDS={body_words}")
    print(f"ABSTRACT_WORDS={abstract_words}")


if __name__ == "__main__":
    build_report()
