from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "final_report" / "REFERENCES_AND_APPENDICES_80_PLUS_DRAFT.docx"
ASSETS = ROOT / "outputs" / "final_report" / "v12_draft_assets"
PROMPTS = ROOT / "prompts" / "v2"
FONT = "Times New Roman"
GRAY = "EFEFEF"
PALE = "F7F7F7"
RED = "A00000"


REFERENCES = [
    "McDonald, A.M., Cranor, L.F.: The cost of reading privacy policies. I/S: A Journal of Law and Policy for the Information Society 4(3), 543-568 (2008).",
    "Amos, R., Acar, G., Lucherini, E., Kshirsagar, M., Narayanan, A., Mayer, J.: Privacy policies over time: Curation and analysis of a million-document dataset. In: Proceedings of the Web Conference 2021, pp. 2165-2176 (2021). https://doi.org/10.1145/3442381.3450048",
    "Information Commissioner's Office: How should we draft our privacy information? https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/individual-rights/the-right-to-be-informed/how-should-we-draft-our-privacy-information/, last accessed 2026/08/23.",
    "Wilson, S., et al.: The creation and analysis of a website privacy policy corpus. In: Proceedings of the 54th Annual Meeting of the Association for Computational Linguistics, pp. 1330-1340 (2016). https://doi.org/10.18653/v1/P16-1126",
    "Harkous, H., Fawaz, K., Lebret, R., Schaub, F., Shin, K.G., Aberer, K.: Polisis: Automated analysis and presentation of privacy policies using deep learning. In: 27th USENIX Security Symposium, pp. 531-548 (2018).",
    "Ravichander, A., Black, A.W., Wilson, S., Norton, T., Sadeh, N.: Question answering for privacy policies: Combining computational and legal perspectives. In: Proceedings of EMNLP-IJCNLP 2019, pp. 4947-4958 (2019). https://doi.org/10.18653/v1/D19-1500",
    "Ravichander, A., et al.: Breaking down walls of text: How can NLP benefit consumer privacy? In: Proceedings of the 59th Annual Meeting of the Association for Computational Linguistics, pp. 4125-4140 (2021). https://doi.org/10.18653/v1/2021.acl-long.319",
    "Singh, J., Fazili, S., Jain, R., Akhtar, M.S.: EROS: Entity-driven controlled policy document summarization. In: Proceedings of LREC-COLING 2024, pp. 6236-6246 (2024). https://aclanthology.org/2024.lrec-main.551/",
    "T.y.s.s, S., Weiss, C., Grabmair, M.: LexSumm and LexT5: Benchmarking and modeling legal summarization tasks in English. In: Proceedings of the Natural Legal Language Processing Workshop 2024, pp. 381-403 (2024). https://doi.org/10.18653/v1/2024.nllp-1.35",
    "Heddaya, M., MacMillan, K., Mei, H., Tan, C., Malani, A.: CaseSumm: A large-scale dataset for long-context summarization from U.S. Supreme Court opinions. In: Findings of NAACL 2025, pp. 1917-1942 (2025). https://doi.org/10.18653/v1/2025.findings-naacl.102",
    "Fabbri, A.R., Kryscinski, W., McCann, B., Xiong, C., Socher, R., Radev, D.: SummEval: Re-evaluating summarization evaluation. Transactions of the Association for Computational Linguistics 9, 391-409 (2021). https://doi.org/10.1162/tacl_a_00373",
    "Krishna, K., et al.: LongEval: Guidelines for human evaluation of faithfulness in long-form summarization. In: Proceedings of EACL 2023, pp. 1650-1669 (2023). https://doi.org/10.18653/v1/2023.eacl-main.121",
    "Liu, Y., Iter, D., Xu, Y., Wang, S., Xu, R., Zhu, C.: G-Eval: NLG evaluation using GPT-4 with better human alignment. In: Proceedings of EMNLP 2023, pp. 2511-2522 (2023). https://doi.org/10.18653/v1/2023.emnlp-main.153",
    "Tang, L., Laban, P., Durrett, G.: MiniCheck: Efficient fact-checking of LLMs on grounding documents. In: Proceedings of EMNLP 2024, pp. 8818-8847 (2024). https://doi.org/10.18653/v1/2024.emnlp-main.499",
    "Dai, X., Karimi, S., Fang, B.: A critical look at meta-evaluating summarisation evaluation metrics. In: Findings of EMNLP 2024, pp. 14795-14808 (2024). https://doi.org/10.18653/v1/2024.findings-emnlp.869",
    "Subbiah, M., Ladhak, F., Mishra, A., Adams, G.T., Chilton, L., McKeown, K.: STORYSUMM: Evaluating faithfulness in story summarization. In: Proceedings of EMNLP 2024, pp. 9988-10005 (2024). https://doi.org/10.18653/v1/2024.emnlp-main.557",
]


PROMPT_SPECS = [
    ("C.1 Basic direct prompt", "basic_direct_v1.2.txt", "88bae524c263eaf00d7f64c17698f966ce932962a0be1c622ee446f766c64211"),
    ("C.2 Basic role-guided prompt", "basic_role_guided_v1.2.txt", "6f8d8dd96225f6845d8f9229213c8ee0f316b255c7fadd4f26c419ed582dbc4c"),
    ("C.3 Basic structured prompt", "basic_structured_v1.2.txt", "415673d27b208d32d3ea46d5c590bfbb5a13b94b6fff51abf69e963f681afe6b"),
    ("C.4 Safety-focused direct prompt", "safety_focused_direct_v1.2.txt", "b09d244c5679b1c424463ffa62078260339415ab00846b061d3e8f371fd687ac"),
    ("C.5 Safety-focused role-guided prompt", "safety_focused_role_guided_v1.2.txt", "799bc774026fe5e6eba6328469b1c449ea53038b876084f2fe79d71dd0765b4d"),
    ("C.6 Safety-focused structured prompt", "safety_focused_structured_v1.2.txt", "b7827e153256f07bf009c2fe25a93387aba2c704d2357c89ddc8add49fd77334"),
]


def set_font(run, size=10, bold=False, color=None, name=FONT):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.6)
    section.bottom_margin = Cm(4.6)
    section.left_margin = Cm(4.4)
    section.right_margin = Cm(4.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0

    for name, size, before, after in (("Heading 1", 12, 12, 7), ("Heading 2", 10, 8, 4)):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    caption = doc.styles["Caption"]
    caption.font.name = FONT
    caption._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    caption.font.size = Pt(9)
    caption.font.italic = False
    caption.font.color.rgb = None
    caption.paragraph_format.space_before = Pt(3)
    caption.paragraph_format.space_after = Pt(6)

    reference = doc.styles.add_style("LNCS Reference", 1)
    reference.base_style = normal
    reference.font.name = FONT
    reference._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    reference._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    reference.font.size = Pt(8.5)
    reference.paragraph_format.left_indent = Cm(0.65)
    reference.paragraph_format.first_line_indent = Cm(-0.65)
    reference.paragraph_format.space_after = Pt(3)

    pending = doc.styles.add_style("Pending Result", 1)
    pending.base_style = normal
    pending.font.name = FONT
    pending._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    pending._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    pending.font.size = Pt(9.5)
    pending.font.bold = True
    pending.font.color.rgb = RGBColor.from_string(RED)
    pending.paragraph_format.space_after = Pt(5)

    add_page_number(section.footer.paragraphs[0])


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=80):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row_index, row in enumerate(table.rows):
        tr_pr = row._tr.get_or_add_trPr()
        tr_pr.append(OxmlElement("w:cantSplit"))
        if row_index == 0:
            repeat = OxmlElement("w:tblHeader")
            repeat.set(qn("w:val"), "true")
            tr_pr.append(repeat)
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)


def add_label_table(doc, caption, headers, rows, widths):
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(caption)
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, value in enumerate(headers):
        shade(table.rows[0].cells[i], GRAY)
        p = table.rows[0].cells[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(value), 8, True)
    for values in rows:
        row = table.add_row()
        for i, value in enumerate(values):
            p = row.cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            set_font(p.add_run(value), 8)
    set_table_geometry(table, widths)


def add_figure(doc, filename, caption, description, width_cm):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(str(ASSETS / filename), width=Cm(width_cm))
    shape._inline.docPr.set("title", caption.split(".")[0])
    shape._inline.docPr.set("descr", description)
    cap = doc.add_paragraph(style="Caption")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.add_run(caption)


def add_prompt_block(doc, title, filename, digest):
    doc.add_heading(title, level=2)
    note = doc.add_paragraph(style="Normal")
    set_font(note.add_run(f"File: {filename}. SHA-256: {digest}"), 8)
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    cell = table.cell(0, 0)
    shade(cell, PALE)
    text = (PROMPTS / filename).read_text(encoding="utf-8").strip()
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    set_font(p.add_run(text), 7.8, name="Consolas")
    set_table_geometry(table, [7200])


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    doc.add_heading("References", level=1)
    for idx, reference in enumerate(REFERENCES, start=1):
        p = doc.add_paragraph(style="LNCS Reference")
        p.add_run(f"[{idx}] {reference}")

    doc.add_page_break()
    doc.add_heading("Appendix A: Code Repository and Reproducibility", level=1)
    pending = doc.add_paragraph(style="Pending Result")
    pending.add_run("[INSERT BEFORE SUBMISSION] Public GitHub repository URL, release/tag and commit hash matching the submitted report.")
    add_body(doc, "The repository should contain the source code, frozen configuration, prompt files, preprocessing and validation commands, anonymised metadata, analysis scripts, MiniCheck notebook, Streamlit PolicyLens prototype and a reproducibility README. API keys, restricted blinding keys and any researcher-identifying files must remain excluded. Where provider endpoints or model versions later change, the submitted release and logged response metadata preserve the exact experimental record.")

    doc.add_page_break()
    doc.add_heading("Appendix B: Modelling Diagrams", level=1)
    add_body(doc, "The diagrams document the experimental design and the two directions of source-grounded verification described in Section 3.")
    add_figure(doc, "figure_B1_v12_pipeline.png", "Fig. B1. Experimental and evaluation pipeline.", "Pipeline from three frozen privacy policies through 54 matched summaries, automated and researcher evaluation, analysis and the PolicyLens prototype.", 12.0)
    doc.add_page_break()
    add_figure(doc, "figure_B2_v12_faithfulness.png", "Fig. B2. Summary-to-source faithfulness workflow.", "Workflow in which each of 1,393 frozen summary sentences is checked against the complete source policy by Gemini, MiniCheck and three blinded researchers, followed by evidence review and agreement analysis.", 12.0)
    doc.add_page_break()
    add_figure(doc, "figure_B3_v12_coverage.png", "Fig. B3. Source-to-summary coverage workflow.", "Workflow in which researchers identify important policy information source-first, adjudicate and freeze a shared unit set, then score each unit against summaries before human-Gemini agreement analysis.", 12.0)

    doc.add_page_break()
    doc.add_heading("Appendix C: Frozen Generation Prompts", level=1)
    add_body(doc, "The six v1.2 prompts preserve the same audience, output format and source placeholder. The safety-focused set adds a common safeguard block; role guidance and internal planning are varied only by strategy. Hashes permit exact verification of the files used for all 54 final summaries.")
    for index, spec in enumerate(PROMPT_SPECS):
        if index in {2, 4, 5}:
            doc.add_page_break()
        add_prompt_block(doc, *spec)

    doc.add_page_break()
    doc.add_heading("Appendix D: Evaluation Codebooks", level=1)
    doc.add_heading("D.1 Summary-to-source faithfulness", level=2)
    add_label_table(
        doc,
        "Table D1. Four-class faithfulness codebook used by Gemini and the researchers.",
        ["Label", "Operational definition", "Failure type"],
        [
            ["Supported", "Every material part of the sentence is supported by the policy.", "None"],
            ["Partially supported", "Core content has source support, but a material part is absent, overstated, imprecise or meaning-altering.", "Usually distortion; combined where unsupported content is also present"],
            ["Unsupported", "The policy provides no adequate evidence for the material statement.", "Hallucination"],
            ["Contradicted", "The sentence conflicts materially with what the policy states.", "Distortion or combined, according to the stated reason"],
        ],
        [1450, 4000, 1750],
    )
    add_body(doc, "Each judgement should include exact source evidence where applicable and a concise reason. MiniCheck remains a separate binary support prediction and must not be presented as producing the four-class labels or measuring omission.")

    doc.add_heading("D.2 Source-to-summary coverage", level=2)
    add_label_table(
        doc,
        "Table D2. Three-class coverage codebook for the shared source units.",
        ["Label", "Score", "Operational definition"],
        [
            ["Covered", "2", "The summary preserves the unit's essential meaning and all material qualifications accurately."],
            ["Partially covered", "1", "The core meaning appears, but a meaningful qualifier, scope, condition or detail is lost."],
            ["Not covered", "0", "The unit is absent, materially distorted or contradicted in the summary."],
        ],
        [1700, 700, 4800],
    )
    add_body(doc, "Important information is identified independently from each complete policy before summaries are examined. No fixed privacy taxonomy is used in the final v1.2 protocol, and a topic absent from the source cannot be counted as an omission. Strict coverage is Covered divided by all frozen units; weighted coverage is (Covered + 0.5 x Partially covered) divided by all frozen units.")

    doc.core_properties.title = "References and Appendices - Verifying LLM-Generated Plain-English Summaries of Privacy Policies"
    doc.core_properties.subject = "Standalone MSc AI dissertation references and appendices"
    doc.core_properties.author = "MSc AI project team"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"REFERENCES={len(REFERENCES)} PROMPTS={len(PROMPT_SPECS)} FIGURES=3")


if __name__ == "__main__":
    build()
