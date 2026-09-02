from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "final_report" / "INTRODUCTION_80_PLUS_DRAFT.docx"
FONT = "Times New Roman"


def set_font(run, size=10, bold=False, italic=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.keep_together = False
    p.add_run(text)
    return p


def add_rq(doc, label, text):
    p = doc.add_paragraph(style="Research Question")
    r = p.add_run(f"{label}. ")
    set_font(r, 10, bold=True)
    p.add_run(text)


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.6)
    section.bottom_margin = Cm(4.6)
    section.left_margin = Cm(4.4)
    section.right_margin = Cm(4.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    h1.font.name = FONT
    h1._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h1._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h1.font.size = Pt(12)
    h1.font.bold = True
    h1.font.color.rgb = None
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(7)
    h1.paragraph_format.keep_with_next = True

    h2 = doc.styles["Heading 2"]
    h2.font.name = FONT
    h2._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    h2._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    h2.font.size = Pt(10)
    h2.font.bold = True
    h2.font.color.rgb = None
    h2.paragraph_format.space_before = Pt(8)
    h2.paragraph_format.space_after = Pt(4)
    h2.paragraph_format.keep_with_next = True

    rq = doc.styles.add_style("Research Question", 1)
    rq.base_style = normal
    rq.font.name = FONT
    rq._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    rq._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    rq.font.size = Pt(10)
    rq.paragraph_format.left_indent = Cm(0.45)
    rq.paragraph_format.first_line_indent = Cm(-0.45)
    rq.paragraph_format.space_after = Pt(4)
    rq.paragraph_format.keep_together = True

    add_page_number(section.footer.paragraphs[0])


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    doc.add_heading("1 Introduction", level=1)

    add_body(
        doc,
        "Privacy policies explain how organisations collect, use, share, retain and protect personal information. They influence decisions about creating accounts, providing data, using services and exercising privacy rights. Yet formal availability does not ensure practical accessibility: notices are often long and dependent on legal or technical vocabulary. McDonald and Cranor quantified the collective burden of reading online privacy policies [1], while a later million-document study found that policies had become longer and more difficult to read [2]. Regulatory guidance therefore stresses concise, transparent, intelligible and clear privacy information [3]. Information may consequently be disclosed while remaining difficult for an ordinary reader to interpret and use."
    )
    add_body(
        doc,
        "Large language models (LLMs) may address this problem by condensing lengthy documents and re-expressing complex wording. However, fluency is not evidence of reliability. A model may introduce unsupported information, materially change a condition or qualification, or omit information that a reader needs. This study terms these failures hallucination, distortion and omission, respectively. The challenge is therefore to improve accessibility without misrepresenting the source. Greater compression and simpler wording may help a reader, but they may also remove context, exceptions and legally significant distinctions such as 'may', 'will' or 'only'."
    )

    add_body(
        doc,
        "Previous research has shown that privacy policies can be classified, queried and summarised, while legal-summarisation studies demonstrate that LLMs can generate fluent condensed accounts of complex documents [4-10]. Evaluation research has also developed automated and human protocols for readability and factual consistency [11-16]. Nevertheless, a summary may contain apparently supported sentences yet remain inadequate because it excludes an important practice, limitation or choice. It may also retain a topic while altering its certainty, actor, scope or exception. Readability alone cannot expose these problems, and checking only generated statements cannot establish that important source content was retained."
    )
    add_body(
        doc,
        "A further issue is whether automated evaluators are credible in this domain. Their scores may be efficient, but agreement between automated systems is not human validation, and aggregate scores may conceal rare failures. The research gap is therefore a lack of controlled, evidence-linked evaluation that treats accessibility, factual meaning and information retention as distinct but related qualities and compares automated decisions with independent researcher judgements. This study addresses the gap bidirectionally: summary-to-source assessment checks hallucination and distortion, while source-to-summary assessment checks important information for omission."
    )

    doc.add_heading("1.1 Aim, Contribution and Research Questions", level=2)
    add_body(
        doc,
        "This study investigates whether selected general-purpose LLMs can generate readable and concise plain-English privacy-policy summaries while preserving source meaning and important information. It uses three complete policies representing short, medium and long/complex documents. GPT, Llama and Mistral are tested under direct, role-guided and structured prompting with matched basic and safety-focused prompts, producing 54 summaries and 27 intervention pairs. Evaluation combines readability and compression, Gemini and MiniCheck verification, and blinded judgements from three student researchers. The PolicyLens prototype exposes labels, evidence and reasons for inspection."
    )
    add_body(
        doc,
        "The novelty is not simply using LLMs for privacy-policy summarisation, which has precedents. The principal contribution is an evidence-linked framework that evaluates the transformation in both directions and connects decisions to an interactive prototype. The matched intervention tests whether safeguards reduce failures without sacrificing accessibility. The restricted corpus supports detailed comparison but not claims about all policies, models or readers. Researcher assessment evaluates source alignment rather than proving end-user comprehension."
    )

    add_rq(doc, "RQ1", "To what extent can the selected LLMs generate privacy-policy summaries that are readable, concise and aligned with the source?")
    add_rq(doc, "RQ2", "How effectively does safety-focused prompting reduce hallucination, distortion and omission while maintaining readability and conciseness?")
    add_rq(doc, "RQ3", "To what extent do automated evaluator judgements align with researcher judgements when identifying hallucination, distortion and omission?")
    add_body(
        doc,
        "Section 2 reviews relevant research and derives the gap. Section 3 describes the corpus, prompting intervention, verification, researcher protocol and prototype. Section 4 presents and discusses the quantitative and qualitative results. Section 5 concludes the study and proposes future work; supporting materials appear in the appendices."
    )

    text = "\n".join(p.text for p in doc.paragraphs)
    count = len(re.findall(r"\b[\w'-]+\b", text))
    doc.core_properties.title = "Introduction - Verifying LLM-Generated Plain-English Summaries of Privacy Policies"
    doc.core_properties.subject = "Standalone MSc AI dissertation introduction draft"
    doc.core_properties.author = "MSc AI Project Team"
    doc.core_properties.keywords = "privacy policy, LLM, plain English, verification, dissertation"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"WORDS={count}")


if __name__ == "__main__":
    build()
