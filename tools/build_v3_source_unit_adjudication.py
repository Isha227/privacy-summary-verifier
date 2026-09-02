"""Create the v3 source-unit adjudication workbook from valid researcher ratings."""

from __future__ import annotations

import csv
import json
import shutil
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data/v2_prompt_intervention_v1_2/source_unit_method_validation/v3_researcher_comparison"
RATINGS = BASE / "validation/v3_researcher_ratings_all.csv"
SOURCE_WORKBOOK = BASE / "Researcher_A1_Package/source_unit_comparison_v3_A1_COMPLETED.xlsx"
OUT = BASE / "adjudication"

NAVY = "163E68"
BLUE = "DDEBF7"
YELLOW = "FFF2CC"
GREY = "F5F7FA"
WHITE = "FFFFFF"
GRID = Side(style="thin", color="D8DEE8")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def unit_lookup() -> dict[str, dict[str, str]]:
    workbook = load_workbook(SOURCE_WORKBOOK, read_only=True, data_only=True)
    result: dict[str, dict[str, str]] = {}
    for sheet_name in ("Set A units", "Set B units"):
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if not row[1]:
                continue
            result[str(row[1]).strip()] = {
                "policy": str(row[0]).strip(),
                "unit_id": str(row[1]).strip(),
                "information": str(row[2] or "").strip(),
                "evidence": str(row[3] or "").strip(),
                "qualifiers": str(row[4] or "").strip(),
                "importance": str(row[5] or "").strip(),
            }
    return result


def apply_title(sheet, text: str, last_col: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = sheet.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos Display", size=14, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center", shrink_to_fit=True)
    sheet.row_dimensions[1].height = 30


def style_header(sheet, row: int, headers: list[str]) -> None:
    for column, value in enumerate(headers, 1):
        cell = sheet.cell(row, column, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=9, bold=True, color=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=GRID)
    sheet.row_dimensions[row].height = 32


def add_dropdown(sheet, target: str, values: list[str]) -> None:
    validation = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    validation.error = "Choose one permitted final label."
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add(target)


def set_print_layout(sheet, landscape: bool = False) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def create_start(workbook: Workbook, candidate_count: int, reference_count: int) -> None:
    sheet = workbook.active
    sheet.title = "START HERE"
    sheet.sheet_view.showGridLines = False
    apply_title(sheet, "V3 source-unit adjudication", 6)
    items = [
        (3, "Purpose", f"Resolve {candidate_count + reference_count} non-unanimous decisions: {candidate_count} Set A candidate decisions and {reference_count} Set B completeness decisions."),
        (5, "Adjudication meaning", "Adjudication is an evidence-based final ruling, not averaging, majority voting or choosing a middle label. Re-read the unit, original evidence and all three reasons before deciding."),
        (7, "Matching IDs", "The final matching IDs must identify every unit required to support the final decision. Several narrow units may collectively match one broad unit. Use semicolons between multiple IDs."),
        (9, "Final labels", "Use the same decision definitions as the independent phase. Preserve a partial label when important actors, data types, purposes, recipients, conditions, exceptions, quantities, time periods or certainty are missing."),
        (11, "Discussion", "The three researchers may now discuss these disputed rows. Record the final evidence-based label and reason, including why the alternative interpretation was rejected."),
        (13, "Completion", "Complete Final matching IDs where the label requires a match, Final label and Final reason for every row. The Completion Check must show YES before submission."),
        (15, "Save as", "source_unit_comparison_v3_ADJUDICATED.xlsx"),
    ]
    for row, heading, value in items:
        sheet.cell(row, 1, heading).font = Font(name="Aptos", size=11, bold=True, color=NAVY)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=6)
        cell = sheet.cell(row, 2, value)
        cell.font = Font(name="Aptos", size=10, color="182230")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        cell.fill = PatternFill("solid", fgColor=GREY)
        sheet.row_dimensions[row].height = 44
    for col, width in enumerate([24, 22, 22, 22, 22, 22], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    set_print_layout(sheet)


def create_disagreement_sheet(workbook: Workbook, name: str, groups: list[list[dict[str, str]]], units: dict[str, dict[str, str]], labels: list[str], expected_prefix: str) -> None:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    headers = [
        "Policy", "Unit ID", "Important information", "Source evidence", "Material qualifiers",
        "A1 decision", "A2 decision", "A3 decision",
        "Final matching IDs", "Final label", "Final reason",
    ]
    apply_title(sheet, name, len(headers))
    style_header(sheet, 3, headers)
    for row_number, group in enumerate(groups, 4):
        by_researcher = {row["researcher"]: row for row in group}
        unit = units[group[0]["unit_id"]]
        values = [unit["policy"], unit["unit_id"], unit["information"], unit["evidence"], unit["qualifiers"]]
        for researcher in ("A1", "A2", "A3"):
            rating = by_researcher[researcher]
            values.append(
                f"Label: {rating['label']}\n"
                f"Matching IDs: {rating['matching_ids'] or 'None'}\n"
                f"Confidence: {rating['confidence']}\n"
                f"Reason: {rating['reason']}"
            )
        values.extend(["", "", ""])
        for column, value in enumerate(values, 1):
            cell = sheet.cell(row_number, column, value)
            cell.font = Font(name="Aptos", size=9, color="182230")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.border = Border(bottom=GRID)
            if column >= 9:
                cell.fill = PatternFill("solid", fgColor=YELLOW)
        sheet.row_dimensions[row_number].height = 126
    end = len(groups) + 3
    add_dropdown(sheet, f"J4:J{end}", labels)
    sheet.freeze_panes = "F4"
    sheet.auto_filter.ref = f"A3:K{end}"
    widths = [9, 16, 42, 52, 28, 44, 44, 44, 24, 22, 48]
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    set_print_layout(sheet, landscape=True)


def create_completion(workbook: Workbook, candidate_count: int, reference_count: int) -> None:
    sheet = workbook.create_sheet("Completion Check")
    sheet.sheet_view.showGridLines = False
    apply_title(sheet, "Adjudication completion check", 5)
    style_header(sheet, 3, ["Sheet", "Expected rows", "Final labels", "Final reasons", "Status"])
    specs = [(4, "Candidate disagreements", candidate_count), (5, "Reference disagreements", reference_count)]
    for row, name, expected in specs:
        sheet.cell(row, 1, name)
        sheet.cell(row, 2, expected)
        sheet.cell(row, 3, f"=COUNTA('{name}'!J4:J{expected+3})")
        sheet.cell(row, 4, f"=COUNTA('{name}'!K4:K{expected+3})")
        sheet.cell(row, 5, f'=IF(AND(B{row}=C{row},B{row}=D{row}),"Complete","Incomplete")')
        for col in range(1, 6):
            cell = sheet.cell(row, col)
            cell.font = Font(name="Aptos", size=10, color="182230")
            cell.border = Border(bottom=GRID)
    sheet["A8"] = "Ready to submit"
    sheet["A8"].font = Font(name="Aptos", size=11, bold=True, color=NAVY)
    sheet["B8"] = '=IF(AND(E4="Complete",E5="Complete"),"YES","NO — complete every final label and reason")'
    sheet["B8"].font = Font(name="Aptos", size=11, bold=True)
    sheet["B8"].fill = PatternFill("solid", fgColor=YELLOW)
    sheet.merge_cells("B8:E8")
    for col, width in enumerate([28, 18, 18, 18, 34], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    set_print_layout(sheet)


def main() -> None:
    rows = read_csv(RATINGS)
    grouped: dict[tuple[str, str], list[dict[str, str]]] = defaultdict(list)
    for row in rows:
        grouped[(row["direction"], row["unit_id"])].append(row)
    disagreements = []
    unanimous = []
    for (direction, unit_id), group in sorted(grouped.items()):
        labels = {row["label"] for row in group}
        matches = {row["matching_ids"] for row in group}
        if len(labels) > 1 or len(matches) > 1:
            disagreements.append(group)
        else:
            unanimous.append({
                "direction": direction, "policy_code": group[0]["policy_code"], "unit_id": unit_id,
                "final_matching_ids": group[0]["matching_ids"], "final_label": group[0]["label"],
                "final_reason": "Unanimous independent researcher decision.", "resolution": "Unanimous",
            })
    candidate = [group for group in disagreements if group[0]["direction"] == "candidate"]
    reference = [group for group in disagreements if group[0]["direction"] == "reference"]
    if len(candidate) != 10 or len(reference) != 11 or len(unanimous) != 258:
        raise RuntimeError(f"Unexpected resolution counts: candidate={len(candidate)}, reference={len(reference)}, unanimous={len(unanimous)}")

    OUT.mkdir(parents=True, exist_ok=True)
    source_dir = OUT / "source_policies"
    source_dir.mkdir(exist_ok=True)
    package_sources = BASE / "Researcher_A1_Package/source_policies"
    for source_path in package_sources.glob("*.txt"):
        shutil.copy2(source_path, source_dir / source_path.name)
    with (OUT / "unanimous_resolutions.csv").open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["direction", "policy_code", "unit_id", "final_matching_ids", "final_label", "final_reason", "resolution"])
        writer.writeheader()
        writer.writerows(unanimous)

    units = unit_lookup()
    workbook = Workbook()
    create_start(workbook, len(candidate), len(reference))
    create_disagreement_sheet(workbook, "Candidate disagreements", candidate, units, ["Equivalent", "Partially overlaps", "Valid additional unit", "Invalid / not important"], "B")
    create_disagreement_sheet(workbook, "Reference disagreements", reference, units, ["Fully represented", "Partly represented", "Not represented"], "A")
    create_completion(workbook, len(candidate), len(reference))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    destination = OUT / "source_unit_comparison_v3_FOR_ADJUDICATION.xlsx"
    workbook.save(destination)
    manifest = {
        "status": "ready_for_adjudication", "total_units": 279, "unanimous_units": len(unanimous),
        "adjudication_units": len(disagreements), "candidate_disagreements": len(candidate),
        "reference_disagreements": len(reference), "label_disagreements": sum(len({r['label'] for r in g}) > 1 for g in disagreements),
        "matching_id_disagreements": sum(len({r['matching_ids'] for r in g}) > 1 for g in disagreements),
        "workbook": str(destination),
    }
    (OUT / "ADJUDICATION_MANIFEST.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
