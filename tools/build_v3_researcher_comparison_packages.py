"""Build blinded researcher workbooks for Gemini v3 source-unit method validation."""

from __future__ import annotations

import csv
import hashlib
import json
import shutil
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data/v2_prompt_intervention_v1_2/source_unit_method_validation"
GEMINI = BASE / "gemini_v3/frozen_gemini_source_units.csv"
HUMAN = ROOT / "data/v2_prompt_intervention_v1_2/human_evaluation/adjudication/frozen_source_units_FINAL.csv"
SOURCES = ROOT / "data/pilot/clean"
OUT = BASE / "v3_researcher_comparison"

POLICY_CODES = {"PILOT01": "P01", "PILOT02": "P02", "PILOT03": "P03"}
RESEARCHERS = ("A1", "A2", "A3")

NAVY = "163E68"
BLUE = "DDEBF7"
PALE_BLUE = "EAF2F8"
PALE_YELLOW = "FFF2CC"
PALE_GREEN = "E2F0D9"
PALE_RED = "FCE4D6"
WHITE = "FFFFFF"
GREY = "667085"
LIGHT_GREY = "F5F7FA"
GRID = Side(style="thin", color="D8DEE8")


def read_csv(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def title(sheet, text: str, last_col: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=last_col)
    cell = sheet.cell(1, 1, text)
    cell.fill = PatternFill("solid", fgColor=NAVY)
    cell.font = Font(name="Aptos Display", size=14, bold=True, color=WHITE)
    cell.alignment = Alignment(vertical="center", shrink_to_fit=True)
    sheet.row_dimensions[1].height = 30


def header(sheet, row: int, columns: list[str]) -> None:
    for col, value in enumerate(columns, 1):
        cell = sheet.cell(row, col, value)
        cell.fill = PatternFill("solid", fgColor=BLUE)
        cell.font = Font(name="Aptos", size=10, bold=True, color=NAVY)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        cell.border = Border(bottom=GRID)
    sheet.row_dimensions[row].height = 30


def body_style(cell, editable: bool = False) -> None:
    cell.font = Font(name="Aptos", size=10, color="182230")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = Border(bottom=GRID)
    if editable:
        cell.fill = PatternFill("solid", fgColor=PALE_YELLOW)


def set_widths(sheet, widths: list[float]) -> None:
    for index, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(index)].width = width


def set_print_layout(sheet, landscape: bool = True) -> None:
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape" if landscape else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.autoPageBreaks = False
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.4
    sheet.page_margins.bottom = 0.4


def add_validation(sheet, cell_range: str, values: list[str]) -> None:
    validator = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    validator.error = "Choose one of the permitted labels."
    validator.errorTitle = "Invalid label"
    validator.prompt = "Select a label from the list."
    validator.promptTitle = "Required decision"
    validator.showErrorMessage = True
    validator.showInputMessage = True
    sheet.add_data_validation(validator)
    validator.add(cell_range)


def build_start_sheet(workbook: Workbook, researcher: str) -> None:
    sheet = workbook.active
    sheet.title = "START HERE"
    sheet.sheet_view.showGridLines = False
    title(sheet, f"Source-unit comparison — Researcher {researcher}", 6)
    content = [
        (3, "Purpose", "Independently compare two blinded source-unit sets derived from the same three contemporary privacy policies. Do not try to produce a compromise or guess which set came from a human or an automated system."),
        (5, "Blinding", "Do not open the restricted blinding key, earlier comparison results, summaries, evaluator results or another researcher's workbook. Use only this workbook and the supplied source policies."),
        (7, "Step 1 — A against B", "For every Set A unit, search Set B within the same policy. Record all matching Set B IDs, choose one label and give a concise evidence-based reason."),
        (9, "Set A labels", "Equivalent — all material content is represented in Set B.\nPartially overlaps — only some material content is represented.\nValid additional unit — explicit and important source information is absent from Set B.\nInvalid / not important — unsupported, duplicate without independent meaning, or not materially useful to the intended reader."),
        (12, "Step 2 — B completeness", "For every Set B unit, search Set A within the same policy. Record all matching Set A IDs and judge whether the unit is fully, partly or not represented."),
        (14, "Set B labels", "Fully represented — one or more Set A units collectively preserve all material content.\nPartly represented — some but not all material content is preserved.\nNot represented — no Set A unit preserves the proposition."),
        (17, "Decision rule", "Judge meaning, not shared vocabulary. Several narrow units may collectively match one broad unit, and one broad unit may partly match several detailed units. Check actors, data types, purposes, recipients, conditions, exceptions, quantities, time periods and certainty."),
        (20, "Independence", "Work alone and do not discuss decisions until A1, A2 and A3 have all submitted completed workbooks. Disagreements will be adjudicated afterwards against the original policy evidence."),
        (22, "Completion", "Use the Completion Check sheet. Every comparison row requires a label, a reason and confidence. Matching IDs may be blank only for Valid additional unit, Invalid / not important, or Not represented."),
        (24, "Save as", f"source_unit_comparison_v3_{researcher}_COMPLETED.xlsx"),
    ]
    for row, heading, text in content:
        sheet.cell(row, 1, heading).font = Font(name="Aptos", size=11, bold=True, color=NAVY)
        sheet.merge_cells(start_row=row, start_column=2, end_row=row + (1 if "\n" in text else 0), end_column=6)
        c = sheet.cell(row, 2, text)
        c.font = Font(name="Aptos", size=10, color="182230")
        c.alignment = Alignment(wrap_text=True, vertical="top")
        c.fill = PatternFill("solid", fgColor=LIGHT_GREY)
        c.border = Border(left=GRID)
        sheet.row_dimensions[row].height = 42 if "\n" not in text else 72
    set_widths(sheet, [28, 20, 20, 20, 20, 20])
    set_print_layout(sheet, landscape=False)


def build_lookup_sheet(workbook: Workbook, name: str, rows: list[dict[str, str]], candidate: bool) -> None:
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    headers = ["Policy", "Unit ID", "Important information", "Exact source evidence", "Material qualifiers", "Why important"]
    title(sheet, name, len(headers))
    header(sheet, 3, headers)
    for index, row in enumerate(rows, 4):
        values = [
            POLICY_CODES[row["policy_id"]],
            row["safe_id"],
            row["important_information"],
            row["exact_source_evidence"] if candidate else row["exact_source_quote"],
            row.get("material_qualifiers", ""),
            row["importance_reason"] if candidate else row["why_important_to_user"],
        ]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index, col, value)
            body_style(cell)
        sheet.row_dimensions[index].height = 54
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:F{len(rows) + 3}"
    set_widths(sheet, [10, 15, 48, 66, 34, 42])
    set_print_layout(sheet)


def build_candidate_comparison(workbook: Workbook, rows: list[dict[str, str]]) -> None:
    sheet = workbook.create_sheet("1 - Set A against B")
    sheet.sheet_view.showGridLines = False
    headers = [
        "Policy", "Set A ID", "Set A important information", "Source evidence", "Material qualifiers",
        "Matching Set B ID(s)", "Decision label", "Reason", "Confidence",
    ]
    title(sheet, "Step 1 — Compare every Set A unit with Set B", len(headers))
    header(sheet, 3, headers)
    for index, row in enumerate(rows, 4):
        values = [
            POLICY_CODES[row["policy_id"]], row["safe_id"], row["important_information"],
            row["exact_source_evidence"], row.get("material_qualifiers", ""), "", "", "", "",
        ]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index, col, value)
            body_style(cell, editable=col >= 6)
        sheet.row_dimensions[index].height = 66
    end = len(rows) + 3
    add_validation(sheet, f"G4:G{end}", ["Equivalent", "Partially overlaps", "Valid additional unit", "Invalid / not important"])
    add_validation(sheet, f"I4:I{end}", ["High", "Medium", "Low"])
    sheet.conditional_formatting.add(f"G4:G{end}", FormulaRule(formula=[f'G4="Valid additional unit"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
    sheet.conditional_formatting.add(f"G4:G{end}", FormulaRule(formula=[f'G4="Invalid / not important"'], fill=PatternFill("solid", fgColor=PALE_RED)))
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:I{end}"
    set_widths(sheet, [9, 14, 48, 60, 30, 24, 22, 48, 13])
    set_print_layout(sheet)


def build_reference_completeness(workbook: Workbook, rows: list[dict[str, str]]) -> None:
    sheet = workbook.create_sheet("2 - Set B completeness")
    sheet.sheet_view.showGridLines = False
    headers = [
        "Policy", "Set B ID", "Set B important information", "Source evidence", "Material qualifiers",
        "Matching Set A ID(s)", "Representation label", "Reason", "Confidence",
    ]
    title(sheet, "Step 2 — Check whether every Set B unit appears in Set A", len(headers))
    header(sheet, 3, headers)
    for index, row in enumerate(rows, 4):
        values = [
            POLICY_CODES[row["policy_id"]], row["safe_id"], row["important_information"],
            row["exact_source_quote"], row.get("material_qualifiers", ""), "", "", "", "",
        ]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(index, col, value)
            body_style(cell, editable=col >= 6)
        sheet.row_dimensions[index].height = 66
    end = len(rows) + 3
    add_validation(sheet, f"G4:G{end}", ["Fully represented", "Partly represented", "Not represented"])
    add_validation(sheet, f"I4:I{end}", ["High", "Medium", "Low"])
    sheet.conditional_formatting.add(f"G4:G{end}", FormulaRule(formula=[f'G4="Fully represented"'], fill=PatternFill("solid", fgColor=PALE_GREEN)))
    sheet.conditional_formatting.add(f"G4:G{end}", FormulaRule(formula=[f'G4="Not represented"'], fill=PatternFill("solid", fgColor=PALE_RED)))
    sheet.freeze_panes = "C4"
    sheet.auto_filter.ref = f"A3:I{end}"
    set_widths(sheet, [9, 14, 48, 60, 30, 24, 22, 48, 13])
    set_print_layout(sheet)


def build_completion_sheet(workbook: Workbook, candidate_count: int, reference_count: int) -> None:
    sheet = workbook.create_sheet("Completion Check")
    sheet.sheet_view.showGridLines = False
    title(sheet, "Completion check", 5)
    header(sheet, 3, ["Task", "Expected rows", "Labels completed", "Reasons completed", "Confidence completed"])
    rows = [
        (4, "Set A against Set B", candidate_count, f"=COUNTA('1 - Set A against B'!G4:G{candidate_count+3})", f"=COUNTA('1 - Set A against B'!H4:H{candidate_count+3})", f"=COUNTA('1 - Set A against B'!I4:I{candidate_count+3})"),
        (5, "Set B completeness", reference_count, f"=COUNTA('2 - Set B completeness'!G4:G{reference_count+3})", f"=COUNTA('2 - Set B completeness'!H4:H{reference_count+3})", f"=COUNTA('2 - Set B completeness'!I4:I{reference_count+3})"),
    ]
    for row_num, label, expected, labels, reasons, confidence in rows:
        values = [label, expected, labels, reasons, confidence]
        for col, value in enumerate(values, 1):
            cell = sheet.cell(row_num, col)
            if isinstance(value, str) and value.startswith("="):
                cell.value = value
            else:
                cell.value = value
            body_style(cell)
    sheet["A8"] = "Ready to submit"
    sheet["A8"].font = Font(name="Aptos", size=11, bold=True, color=NAVY)
    sheet["B8"] = "=IF(AND(B4=C4,B4=D4,B4=E4,B5=C5,B5=D5,B5=E5),\"YES\",\"NO — complete all required cells\")"
    sheet["B8"].font = Font(name="Aptos", size=11, bold=True)
    sheet["B8"].fill = PatternFill("solid", fgColor=PALE_YELLOW)
    sheet.merge_cells("B8:E8")
    set_widths(sheet, [28, 18, 20, 20, 22])
    set_print_layout(sheet, landscape=False)


def build_workbook(researcher: str, candidates: list[dict[str, str]], references: list[dict[str, str]], path: Path) -> None:
    workbook = Workbook()
    build_start_sheet(workbook, researcher)
    build_lookup_sheet(workbook, "Set A units", candidates, candidate=True)
    build_lookup_sheet(workbook, "Set B units", references, candidate=False)
    build_candidate_comparison(workbook, candidates)
    build_reference_completeness(workbook, references)
    build_completion_sheet(workbook, len(candidates), len(references))
    workbook.calculation.fullCalcOnLoad = True
    workbook.calculation.forceFullCalc = True
    workbook.calculation.calcMode = "auto"
    workbook.save(path)


def validate_workbook(path: Path, candidate_count: int, reference_count: int) -> dict[str, object]:
    workbook = load_workbook(path, data_only=False)
    required = ["START HERE", "Set A units", "Set B units", "1 - Set A against B", "2 - Set B completeness", "Completion Check"]
    errors: list[str] = []
    if workbook.sheetnames != required:
        errors.append(f"Unexpected sheets: {workbook.sheetnames}")
    if workbook["Set A units"].max_row != candidate_count + 3:
        errors.append("Set A row count mismatch")
    if workbook["Set B units"].max_row != reference_count + 3:
        errors.append("Set B row count mismatch")
    if workbook["1 - Set A against B"].max_row != candidate_count + 3:
        errors.append("Candidate comparison row count mismatch")
    if workbook["2 - Set B completeness"].max_row != reference_count + 3:
        errors.append("Reference comparison row count mismatch")
    if len(workbook["1 - Set A against B"].data_validations.dataValidation) != 2:
        errors.append("Candidate validations missing")
    if len(workbook["2 - Set B completeness"].data_validations.dataValidation) != 2:
        errors.append("Reference validations missing")
    return {"path": str(path), "sha256": sha256(path), "errors": errors}


def main() -> None:
    candidates = read_csv(GEMINI)
    references = read_csv(HUMAN)
    if len(candidates) != 114 or len(references) != 165:
        raise RuntimeError(f"Unexpected frozen counts: candidates={len(candidates)}, references={len(references)}")

    candidate_rows: list[dict[str, str]] = []
    reference_rows: list[dict[str, str]] = []
    key_rows: list[dict[str, str]] = []
    per_policy_candidate = {policy_id: 0 for policy_id in POLICY_CODES}
    per_policy_reference = {policy_id: 0 for policy_id in POLICY_CODES}
    for row in candidates:
        policy_id = row["policy_id"]
        per_policy_candidate[policy_id] += 1
        safe_id = f"A-{POLICY_CODES[policy_id]}-U{per_policy_candidate[policy_id]:03d}"
        candidate_rows.append({**row, "safe_id": safe_id})
        key_rows.append({"set": "A", "policy_id": policy_id, "safe_id": safe_id, "original_id": row["candidate_unit_id"], "provenance": "Gemini v3"})
    for row in references:
        policy_id = row["policy_id"]
        per_policy_reference[policy_id] += 1
        safe_id = f"B-{POLICY_CODES[policy_id]}-U{per_policy_reference[policy_id]:03d}"
        reference_rows.append({**row, "safe_id": safe_id})
        key_rows.append({"set": "B", "policy_id": policy_id, "safe_id": safe_id, "original_id": row["final_unit_id"], "provenance": "Human-adjudicated reference"})

    OUT.mkdir(parents=True, exist_ok=True)
    restricted = OUT / "restricted"
    restricted.mkdir(exist_ok=True)
    key_path = restricted / "BLINDING_KEY_RESTRICTED.csv"
    with key_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=["set", "policy_id", "safe_id", "original_id", "provenance"])
        writer.writeheader()
        writer.writerows(key_rows)

    audits = []
    for researcher in RESEARCHERS:
        package = OUT / f"Researcher_{researcher}_Package"
        package.mkdir(exist_ok=True)
        workbook_path = package / f"source_unit_comparison_v3_{researcher}.xlsx"
        build_workbook(researcher, candidate_rows, reference_rows, workbook_path)
        source_dir = package / "source_policies"
        source_dir.mkdir(exist_ok=True)
        for policy_id, policy_code in POLICY_CODES.items():
            shutil.copy2(SOURCES / f"{policy_id}.txt", source_dir / f"{policy_code}_source.txt")
        audits.append(validate_workbook(workbook_path, len(candidate_rows), len(reference_rows)))

    manifest = {
        "status": "ready_for_independent_researcher_comparison" if not any(item["errors"] for item in audits) else "failed",
        "researchers": list(RESEARCHERS),
        "candidate_units": len(candidate_rows),
        "reference_units": len(reference_rows),
        "decisions_per_researcher": len(candidate_rows) + len(reference_rows),
        "total_independent_decisions": (len(candidate_rows) + len(reference_rows)) * len(RESEARCHERS),
        "candidate_units_by_policy": per_policy_candidate,
        "reference_units_by_policy": per_policy_reference,
        "gemini_frozen_sha256": sha256(GEMINI),
        "human_reference_sha256": sha256(HUMAN),
        "restricted_key_sha256": sha256(key_path),
        "workbooks": audits,
    }
    (OUT / "PACKAGE_MANIFEST.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    print(json.dumps(manifest, indent=2))


if __name__ == "__main__":
    main()
