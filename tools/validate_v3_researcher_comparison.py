"""Validate and consolidate completed v3 source-unit comparison workbooks."""

from __future__ import annotations

import csv
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "data/v2_prompt_intervention_v1_2/source_unit_method_validation/v3_researcher_comparison"
KEY = BASE / "restricted/BLINDING_KEY_RESTRICTED.csv"
VALIDATION = BASE / "validation"
RESEARCHERS = ("A1", "A2", "A3")
CANDIDATE_LABELS = {"Equivalent", "Partially overlaps", "Valid additional unit", "Invalid / not important"}
REFERENCE_LABELS = {"Fully represented", "Partly represented", "Not represented"}
CONFIDENCE = {"High", "Medium", "Low"}
ID_PATTERN = re.compile(r"[AB]-P0[1-3]-U\d{3}")


def csv_rows(path: Path) -> list[dict[str, str]]:
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def parse_ids(value: object) -> tuple[list[str], str]:
    raw = text(value)
    ids = list(dict.fromkeys(ID_PATTERN.findall(raw)))
    remainder = ID_PATTERN.sub("", raw)
    remainder = re.sub(r"[\s,;|/]+", "", remainder).casefold()
    if remainder in {"", "none", "na", "n.a.", "-"}:
        remainder = ""
    return ids, remainder


def validate_sheet(
    workbook,
    researcher: str,
    sheet_name: str,
    direction: str,
    expected_rows: int,
    labels: set[str],
    expected_set: str,
    key_lookup: dict[str, dict[str, str]],
) -> tuple[list[dict[str, str]], list[str]]:
    sheet = workbook[sheet_name]
    ratings: list[dict[str, str]] = []
    errors: list[str] = []
    for row_number in range(4, expected_rows + 4):
        policy_code = text(sheet.cell(row_number, 1).value)
        unit_id = text(sheet.cell(row_number, 2).value)
        match_ids, remainder = parse_ids(sheet.cell(row_number, 6).value)
        label = text(sheet.cell(row_number, 7).value)
        reason = text(sheet.cell(row_number, 8).value)
        confidence = text(sheet.cell(row_number, 9).value)
        prefix = f"{researcher} {sheet_name} row {row_number} ({unit_id})"
        if not policy_code or not unit_id:
            errors.append(f"{prefix}: missing locked identifier")
            continue
        if label not in labels:
            errors.append(f"{prefix}: invalid or blank label {label!r}")
        if not reason:
            errors.append(f"{prefix}: blank reason")
        if confidence not in CONFIDENCE:
            errors.append(f"{prefix}: invalid or blank confidence {confidence!r}")
        if remainder:
            errors.append(f"{prefix}: unrecognised matching-ID text {remainder!r}")
        for match_id in match_ids:
            target = key_lookup.get(match_id)
            if target is None:
                errors.append(f"{prefix}: unknown matching ID {match_id}")
            else:
                if target["set"] != expected_set:
                    errors.append(f"{prefix}: {match_id} belongs to Set {target['set']}, expected Set {expected_set}")
                if target["policy_id"].replace("PILOT", "P") != policy_code:
                    errors.append(f"{prefix}: cross-policy match {match_id}")
        if direction == "candidate" and label in {"Equivalent", "Partially overlaps"} and not match_ids:
            errors.append(f"{prefix}: {label} requires at least one matching Set B ID")
        if direction == "candidate" and label == "Valid additional unit" and match_ids:
            errors.append(f"{prefix}: Valid additional unit should not have a matching Set B ID")
        if direction == "reference" and label in {"Fully represented", "Partly represented"} and not match_ids:
            errors.append(f"{prefix}: {label} requires at least one matching Set A ID")
        if direction == "reference" and label == "Not represented" and match_ids:
            errors.append(f"{prefix}: Not represented should not have a matching Set A ID")
        ratings.append({
            "researcher": researcher,
            "direction": direction,
            "policy_code": policy_code,
            "unit_id": unit_id,
            "matching_ids": ";".join(match_ids),
            "label": label,
            "reason": reason,
            "confidence": confidence,
            "workbook_row": str(row_number),
        })
    return ratings, errors


def main() -> None:
    key_rows = csv_rows(KEY)
    key_lookup = {row["safe_id"]: row for row in key_rows}
    if len(key_lookup) != 279:
        raise RuntimeError(f"Expected 279 unique blinded IDs, found {len(key_lookup)}")
    all_ratings: list[dict[str, str]] = []
    errors: list[str] = []
    files: dict[str, str] = {}
    for researcher in RESEARCHERS:
        path = BASE / f"Researcher_{researcher}_Package/source_unit_comparison_v3_{researcher}_COMPLETED.xlsx"
        if not path.exists():
            errors.append(f"Missing completed workbook for {researcher}: {path}")
            continue
        files[researcher] = str(path)
        workbook = load_workbook(path, read_only=False, data_only=False)
        if workbook.sheetnames != ["START HERE", "Set A units", "Set B units", "1 - Set A against B", "2 - Set B completeness", "Completion Check"]:
            errors.append(f"{researcher}: unexpected sheet list {workbook.sheetnames}")
        candidate, candidate_errors = validate_sheet(
            workbook, researcher, "1 - Set A against B", "candidate", 114,
            CANDIDATE_LABELS, "B", key_lookup,
        )
        reference, reference_errors = validate_sheet(
            workbook, researcher, "2 - Set B completeness", "reference", 165,
            REFERENCE_LABELS, "A", key_lookup,
        )
        all_ratings.extend(candidate + reference)
        errors.extend(candidate_errors + reference_errors)

    VALIDATION.mkdir(parents=True, exist_ok=True)
    ratings_path = VALIDATION / "v3_researcher_ratings_all.csv"
    fields = ["researcher", "direction", "policy_code", "unit_id", "matching_ids", "label", "reason", "confidence", "workbook_row"]
    with ratings_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        writer.writerows(all_ratings)

    per_researcher = {}
    for researcher in RESEARCHERS:
        rows = [row for row in all_ratings if row["researcher"] == researcher]
        per_researcher[researcher] = {
            "rows": len(rows),
            "candidate_rows": sum(row["direction"] == "candidate" for row in rows),
            "reference_rows": sum(row["direction"] == "reference" for row in rows),
            "labels": dict(sorted(Counter(row["label"] for row in rows).items())),
            "confidence": dict(sorted(Counter(row["confidence"] for row in rows).items())),
        }
    report = {
        "status": "valid" if not errors and len(all_ratings) == 837 else "invalid",
        "expected_total_ratings": 837,
        "actual_total_ratings": len(all_ratings),
        "errors": errors,
        "completed_workbooks": files,
        "per_researcher": per_researcher,
        "ratings_csv": str(ratings_path),
    }
    report_path = VALIDATION / "v3_researcher_validation_report.json"
    report_path.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(json.dumps(report, indent=2))


if __name__ == "__main__":
    main()
