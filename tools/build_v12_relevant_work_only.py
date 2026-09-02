from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "final_report" / "RELEVANT_WORK_80_PLUS_DRAFT.docx"
FONT = "Times New Roman"


def set_font(run, size=10, bold=False):
    run.font.name = FONT
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), FONT)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), FONT)
    run.font.size = Pt(size)
    run.font.bold = bold


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, 8)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, end])


def configure(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(4.6)
    section.bottom_margin = Cm(4.6)
    section.left_margin = Cm(4.4)
    section.right_margin = Cm(4.4)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.5)

    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
    normal.font.size = Pt(10)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(4)
    normal.paragraph_format.line_spacing = 1.0
    normal.paragraph_format.widow_control = True

    for name, size, before, after in (
        ("Heading 1", 12, 12, 7),
        ("Heading 2", 10, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = FONT
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = None
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    add_page_number(section.footer.paragraphs[0])


def add_body(doc, text):
    p = doc.add_paragraph(style="Normal")
    p.add_run(text)
    return p


def build():
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure(doc)

    doc.add_heading("2 Relevant Work", level=1)

    doc.add_heading("2.1 Privacy-Policy Accessibility and Summarisation", level=2)
    add_body(
        doc,
        "Research on privacy notices has consistently identified a gap between formal disclosure and usable information. McDonald and Cranor estimated the substantial time required for individuals to read privacy policies [1], while Amos et al.'s million-document analysis found that policies had grown longer and less readable over time [2]. The Information Commissioner's Office accordingly recommends concise wording, short sentences and information that is easy to locate, while also advising organisations to test notices with intended users [3]. These findings justify plain-language transformation, but they also show why formula-based readability cannot be treated as evidence of comprehension."
    )
    add_body(
        doc,
        "Earlier privacy NLP largely improved access by structuring or retrieving policy content. OPP-115 supplied expert annotations of privacy practices across 115 policies [4], enabling comparable categories and supervised analysis. Polisis extended this approach through automated classification and scalable presentation [5], whereas PrivacyQA framed access as answering questions grounded in policy text [6]. Their strengths are traceability and targeted navigation; their limitation for the present problem is that classification and retrieval do not produce a complete plain-English account. A fixed taxonomy may also privilege predefined practices, while question answering depends on what a user knows to ask."
    )
    add_body(
        doc,
        "Consumer-privacy summarisation has used extractive selection to surface relevant or risky sentences, reducing invention because wording remains close to the source [7]. However, extraction may retain the vocabulary and syntax that made the policy difficult initially. EROS instead used entity-driven controlled abstractive summarisation and reported improvements in readability and information content [8]. Abstraction offers more natural explanation, but its freedom to combine and rephrase material increases the need to verify whether meaning and important breadth have been preserved. These approaches therefore establish useful foundations without resolving full-policy, plain-English verification."
    )

    doc.add_heading("2.2 LLM Legal Summarisation and Reliability", level=2)
    add_body(
        doc,
        "Evidence from broader legal summarisation reinforces this concern. LexSumm benchmarked legal summarisation across eight English datasets and reported abstraction and faithfulness errors in zero-shot LLM output [9]. CaseSumm evaluated long-context summaries of United States Supreme Court opinions and found that favourable automatic scores could coexist with hallucinations, misrepresented facts and disagreement with legal-expert assessment [10]. Although court opinions differ from consumer privacy notices, both domains contain qualifications, actors and conditions whose alteration can materially change meaning. The studies demonstrate technical capability, but also weaken any assumption that fluent or highly rated output is automatically dependable."
    )
    add_body(
        doc,
        "This project distinguishes three related failures. Hallucination introduces material for which the source supplies no adequate evidence. Distortion retains a source basis but changes certainty, scope, actor, condition, exception or relationship. Omission leaves important source information inadequately represented. The distinction matters because a summary can contain only supported sentences yet still mislead through absence, while a retained topic may still be distorted. Prompting may influence these outcomes without retraining: direct instructions provide a baseline, role guidance emphasises audience, and structured guidance encourages organised selection. Explicit safeguards may discourage assumptions and lost qualifications, but could also lengthen the output or reproduce source complexity. Prompt design should therefore be tested as an intervention rather than presumed to improve every dimension."
    )

    doc.add_heading("2.3 Evaluation Methods and Research Gap", level=2)
    add_body(
        doc,
        "Summary quality is multidimensional. Readability formulae and compression ratios are reproducible and permit matched comparison, but they measure textual features rather than understanding, legal adequacy or usability. SummEval similarly showed that automatic metrics capture different qualities and correspond imperfectly with human judgement [11]. For long-form faithfulness, LongEval found that finer-grained evaluation reduced annotator variance [12], supporting sentence-level decisions over a single whole-summary label. Such decomposition improves inspectability, although sentences may still depend on surrounding context."
    )
    add_body(
        doc,
        "LLM-based evaluators can provide labels, evidence and explanations at scale. G-Eval reported stronger correspondence with human judgements than several earlier automatic measures, while also raising concerns about evaluator bias [13]. MiniCheck offers a smaller document-grounded fact checker with a binary support decision [14], but binary classification cannot directly distinguish partial distortion from complete lack of support. Dai et al. further observed that summarisation metric meta-evaluation remains dominated by news datasets [15], limiting confidence in transfer to legal and privacy texts. STORYSUMM showed that both human protocols and automated measures can miss difficult inconsistencies [16]. These findings justify using multiple evaluators, but not treating agreement between automated systems as human validation."
    )
    add_body(
        doc,
        "A further limitation is directional. Faithfulness asks whether generated statements are supported by the source; it does not establish whether important source information was retained. Coverage asks the reverse question and is therefore necessary to examine omission. Existing work provides strong foundations for privacy analysis, legal summarisation and factuality assessment, but controlled evidence remains limited on whether general-purpose LLMs can improve privacy-policy accessibility while preserving both meaning and important breadth. Evidence is also limited on whether automated judgements align with researchers on identical legal-summary units. This study addresses the gap through matched basic and safety-focused prompts, sentence-level summary-to-source verification, source-first coverage assessment, automated-researcher agreement analysis and an evidence-linked interactive prototype."
    )

    text = "\n".join(p.text for p in doc.paragraphs)
    count = len(re.findall(r"\b[\w'-]+\b", text))
    doc.core_properties.title = "Relevant Work - Verifying LLM-Generated Plain-English Summaries of Privacy Policies"
    doc.core_properties.subject = "Standalone MSc AI dissertation relevant work draft"
    doc.core_properties.author = "MSc AI Project Team"
    doc.core_properties.keywords = "privacy policy, legal summarisation, LLM evaluation, faithfulness, coverage"
    doc.save(OUTPUT)
    print(f"OUTPUT={OUTPUT}")
    print(f"WORDS={count}")


if __name__ == "__main__":
    build()
