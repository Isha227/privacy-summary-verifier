from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parent
RESULTS = ROOT / "outputs" / "main_analysis_master" / "Final_Results_Chapter.docx"
DISCUSSION = ROOT / "outputs" / "main_analysis_master" / "Discussion_Chapter_Draft.docx"
OUT = ROOT / "outputs" / "main_analysis_master" / "Results_Discussion_Conclusion.docx"


def add_page_break(doc):
    p = doc.add_paragraph()
    p.add_run().add_break(WD_BREAK.PAGE)


def copy_discussion_without_packaging_or_references(target, source):
    copying = False
    references = []
    in_references = False
    for child in source.element.body.iterchildren():
        if child.tag == qn("w:sectPr"):
            continue
        text = "".join(child.xpath(".//w:t/text()")) .strip()
        if text == "5.1 Overview of the findings":
            copying = True
        if text == "References":
            in_references = True
            copying = False
            continue
        if in_references:
            references.append(deepcopy(child))
        elif copying:
            target.element.body.insert(-1, deepcopy(child))
    return references


def body(doc, text):
    p = doc.add_paragraph(text)
    p.style = "Normal"
    return p


def heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def configure_heading_levels(doc):
    # Preserve the established visual system while making chapter/subsection
    # hierarchy explicit in the combined manuscript.
    h1 = doc.styles["Heading 1"]
    h1.font.name = "Calibri"
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = RGBColor(46, 116, 181)
    h1.paragraph_format.space_before = Pt(16)
    h1.paragraph_format.space_after = Pt(8)
    h1.paragraph_format.keep_with_next = True
    h2 = doc.styles["Heading 2"]
    h2.font.name = "Calibri"
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = RGBColor(46, 116, 181)
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    h2.paragraph_format.keep_with_next = True


def add_conclusion(doc):
    add_page_break(doc)
    heading(doc, "6 Conclusion", 1)
    body(doc, "This dissertation investigated how selected LLM and prompting strategy affected the faithfulness, readability and conciseness of plain-English privacy-policy summaries. A controlled 3 x 3 within-document experiment generated 270 summaries from 30 contemporary policies using GPT-5.6 Luna, Llama 3.3 70B Instruct Turbo and Mistral Large 3 under zero-shot, role and structured-reasoning prompts. The summaries were assessed through automated readability and compression measures, claim-level Gemini and MiniCheck faithfulness evaluation, and blinded human validation.")
    body(doc, "The principal conclusion is that model choice mattered more consistently than prompt choice, but no model dominated every quality dimension. Mistral generated substantially more readable summaries, although these were also the longest and least compressed. Llama produced the shortest summaries and the strongest MiniCheck support rate, while GPT achieved the strongest Gemini support rate. Human labels aligned more closely with Gemini than with MiniCheck on the validation sample, supporting Gemini as the primary faithfulness indicator for this dataset while retaining MiniCheck as an important independent check.")
    body(doc, "The three prompting strategies did not produce a reliable independent improvement in readability or faithfulness. Structured reasoning generated slightly longer outputs and affected compression, but it did not provide a general quality advantage over the simpler prompts. Some model-prompt interactions were present for Gemini-assessed faithfulness and length-related outcomes, but these were not reproduced for MiniCheck faithfulness or readability. Prompt effects were therefore model-, outcome- and evaluator-dependent rather than universal.")
    body(doc, "These findings answer the main research question by demonstrating that plain-English privacy-policy summarisation is a multi-objective selection problem. A model that produces more readable prose may sacrifice conciseness, while the apparent faithfulness leader may change with the evaluator. It would therefore be misleading to collapse all outcomes into a single league table. In a practical system, model selection should follow the intended priority, and generated claims should remain linked to their source evidence so that unsupported or disputed content can be reviewed.")
    body(doc, "The study contributes a reproducible whole-policy comparison across three LLM families and three controlled prompt strategies, together with a claim-level triangulation framework combining two automated evaluators and human judgement. It also provides evidence that evaluator choice can materially influence conclusions about model quality. This reinforces the value of reporting agreement, disagreement and class prevalence rather than relying on one automated factuality score.")
    body(doc, "The conclusions remain bounded by the purposively selected 30-policy corpus, one retained generation per main condition, specific API model versions, formula-based readability, and a human-labelled subset rather than all claims. Most importantly, the complete main-study information-coverage annotation was not available; high faithfulness cannot establish that important policy information was retained. Future work should complete coverage assessment, repeat generations to measure stochastic variability, test comprehension with non-specialist readers, and evaluate evidence-linked revision or verification workflows.")
    body(doc, "Overall, the results favour a cautious, transparent use of LLMs as aids for producing layered privacy explanations rather than replacements for the authoritative policy. The strongest design is not simply the model with the highest mean score, but a workflow that balances accessibility and compression, verifies factual support, exposes evidence, and preserves human oversight where privacy information may affect users' understanding of their rights and an organisation's data practices.")


def add_references(doc, reference_elements):
    add_page_break(doc)
    heading(doc, "References", 1)
    for element in reference_elements:
        doc.element.body.insert(-1, element)


def update_furniture(doc):
    for section in doc.sections:
        header = section.header.paragraphs[0]
        header.clear()
        header.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = header.add_run("MSc AI Dissertation | Results, Discussion and Conclusion")
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(95, 105, 115)
        footer = section.footer.paragraphs[0]
        footer.clear()
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        run = footer.add_run("Main study")
        run.font.name = "Calibri"
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(95, 105, 115)


def main():
    doc = Document(RESULTS)
    discussion = Document(DISCUSSION)
    configure_heading_levels(doc)
    add_page_break(doc)
    heading(doc, "5 Discussion", 1)
    references = copy_discussion_without_packaging_or_references(doc, discussion)
    add_conclusion(doc)
    add_references(doc, references)
    update_furniture(doc)
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
