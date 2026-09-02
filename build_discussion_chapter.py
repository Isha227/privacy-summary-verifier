from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs" / "main_analysis_master" / "Discussion_Chapter_Draft.docx"

BLUE = RGBColor(46, 116, 181)
DARK = RGBColor(31, 77, 120)
MUTED = RGBColor(95, 105, 115)


def font(run, size=11, bold=False, italic=False, color=None):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Calibri")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = color


def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.style = "Normal"
    return p


def add_heading(doc, text, level=1):
    return doc.add_heading(text, level=level)


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text)
    return p


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tag = OxmlElement("w:tblHeader")
    tag.set(qn("w:val"), "true")
    tr_pr.append(tag)


def set_cell_width(cell, dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(dxa))
    tc_w.set(qn("w:type"), "dxa")


def add_summary_table(doc):
    rows = [
        ("RQ1: Model choice", "Strong effect", "Mistral was most readable but longest; Llama was shortest and MiniCheck-favoured; GPT was Gemini-favoured."),
        ("RQ2: Prompt strategy", "Limited independent effect", "No prompt main effect on readability or either faithfulness judge; small effects on length/compression."),
        ("RQ3: Model x prompt", "Evaluator-dependent", "Interaction for Gemini faithfulness and length/compression, but not for MiniCheck faithfulness or readability."),
    ]
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.style = "Table Grid"
    widths = [2300, 1900, 5160]
    hdr = table.rows[0]
    set_repeat_header(hdr)
    for idx, label in enumerate(("Question", "Overall answer", "Interpretation")):
        set_cell_width(hdr.cells[idx], widths[idx])
        set_cell_shading(hdr.cells[idx], "F2F4F7")
        r = hdr.cells[idx].paragraphs[0].add_run(label)
        font(r, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_width(cells[idx], widths[idx])
            cells[idx].text = value
            for p in cells[idx].paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs:
                    font(r, size=10)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), "9360")
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)


def configure(doc):
    sec = doc.sections[0]
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11)
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    sec.header_distance = sec.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ["List Bullet", "List Number"]:
        st = doc.styles[style_name]
        st.font.name = "Calibri"
        st.font.size = Pt(11)
        st.paragraph_format.left_indent = Inches(0.5)
        st.paragraph_format.first_line_indent = Inches(-0.25)
        st.paragraph_format.space_after = Pt(8)
        st.paragraph_format.line_spacing = 1.167

    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = header.add_run("MSc AI Dissertation | Discussion Chapter Draft")
    font(r, size=9, color=MUTED)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Discussion")
    font(r, size=9, color=MUTED)


def main():
    doc = Document()
    configure(doc)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run("DISCUSSION")
    font(r, size=23, bold=True)
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run("Plain-English privacy-policy summarisation across GPT, Llama and Mistral")
    font(r, size=14, color=MUTED)

    add_heading(doc, "5.1 Overview of the findings", 1)
    add_body(doc, "This study examined whether selected large language models (LLMs), prompting strategies, and their interaction influenced the quality of plain-English privacy-policy summaries. The completed main experiment comprised 270 summaries generated from 30 contemporary policies under a fully crossed 3 x 3 design. The clearest overall finding is that model choice mattered more consistently than prompt choice. However, no model was uniformly best: the models occupied different positions on readability, conciseness and faithfulness, and the apparent faithfulness leader depended on the evaluator used.")
    add_body(doc, "Mistral produced substantially more readable summaries than GPT and Llama, but also the longest and least compressed outputs. Llama produced the shortest summaries and achieved the highest MiniCheck support rate, whereas GPT achieved the highest Gemini support rate. Prompt strategy had modest effects on output length and compression, but it did not independently improve Flesch Reading Ease or faithfulness under either Gemini or MiniCheck. Interactions were present for summary length, compression and Gemini faithfulness, but absent for readability and MiniCheck faithfulness. Consequently, the study does not justify recommending a single model-prompt combination for every use case.")
    add_summary_table(doc)

    add_heading(doc, "5.2 RQ1: Effect of LLM choice", 1)
    add_heading(doc, "Readability and conciseness", 2)
    add_body(doc, "RQ1 asked how the choice of LLM affected faithfulness, information coverage, readability and conciseness. For the outcomes completed in the main study, the evidence demonstrates a strong model effect. Mistral's mean Flesch Reading Ease score was 51.16, compared with 27.87 for GPT and 25.71 for Llama. Pairwise comparisons indicated that these differences were statistically significant. This is practically important because privacy information is expected to be concise, intelligible and written in clear, plain language (Information Commissioner's Office, n.d.-a). Mistral therefore came closest of the three systems to the linguistic accessibility objective measured by Flesch, although a score near 51 still does not establish that the summaries were understandable to all intended readers.")
    add_body(doc, "The readability advantage came with a substantial cost. Mistral summaries averaged 781.81 words, whereas GPT averaged 635.38 and Llama 449.62. The model effect on compression remained significant in the policy-fixed-effects sensitivity analysis, reducing the likelihood that the pattern was caused only by differences among source policies. This exposes the central tension in privacy communication identified by regulatory guidance: a notice must include material information while remaining concise and accessible (Information Commissioner's Office, n.d.-b). Mistral appears to have simplified through fuller explanation rather than through aggressive condensation. Llama did the reverse, delivering the greatest compression but not the easiest prose according to Flesch.")
    add_body(doc, "This distinction also shows why word count and readability should not be treated as interchangeable. Shorter text may reduce reading burden without simplifying sentence structure or vocabulary, while longer text may explain legal concepts more directly. Prior longitudinal work has similarly treated policy length and readability as separate properties (Amos et al., 2022). In practical terms, an organisation seeking a short first-layer notice might favour Llama's compression, while one prioritising accessible explanation might prefer Mistral, subject to a maximum length or post-generation editing stage.")

    add_heading(doc, "Faithfulness", 2)
    add_body(doc, "Model rankings for faithfulness were not stable across evaluators. Gemini assigned GPT the highest strict support rate (97.66%), followed by Mistral (96.29%) and Llama (92.99%). MiniCheck instead favoured Llama (91.90%), followed by GPT (83.52%) and Mistral (81.60%). Human judgements on the hybrid sample agreed exactly with Gemini in 92.73% of cases and agreed on strict binary support in 92.89%, whereas strict agreement with MiniCheck was 81.16%. On this sample, Gemini therefore aligned more closely with the human labels, but that does not establish that Gemini was objectively correct in every disagreement.")
    add_body(doc, "The divergence is methodologically informative rather than merely inconvenient. MiniCheck was designed as an efficient claim-grounding model and was reported to approach GPT-4-level fact-checking performance at much lower cost (Tang, Laban and Durrett, 2024). Nevertheless, factual-consistency metrics are sensitive to their architectures, evidence selection and decision thresholds. Earlier work found that automatic summarisation metrics can disagree when summaries occupy a narrow score range (Bhandari et al., 2020), while recent long-document stress testing reports reduced robustness for information-dense claims and long-range evidence (Mujahid, Wright and Augenstein, 2026). Privacy policies present exactly these conditions: repeated concepts, dispersed qualifications and long source documents.")
    add_body(doc, "Accordingly, the defensible conclusion for RQ1 is not that GPT or Llama was categorically most faithful. Instead, GPT was strongest under the LLM judge, Llama was strongest under the specialised local checker, and the human validation sample supported greater reliance on Gemini for the present dataset. This evaluator-contingent result should be preserved in the dissertation because collapsing the two automated scores into a single average would conceal a substantive measurement issue.")

    add_heading(doc, "Information coverage", 2)
    add_body(doc, "The approved RQ1 also named information coverage. A source-unit coverage framework was developed and piloted, but a complete main-study coverage annotation was not produced for the 270 summaries. Faithfulness cannot substitute for coverage: a summary may contain only supported claims while omitting important source information. Therefore, the completed study can answer RQ1 for faithfulness, readability and conciseness, but cannot make a full-corpus claim about comparative coverage. This is a scope limitation and should be stated explicitly in the final dissertation rather than treating high claim support as evidence of completeness.")

    add_heading(doc, "5.3 RQ2: Effect of prompting strategy", 1)
    add_body(doc, "RQ2 asked whether zero-shot, role and structured-reasoning prompts affected summary quality. The overall evidence provides little support for an independent prompt advantage. Prompt strategy was not significant for Flesch Reading Ease, Gemini faithfulness or MiniCheck faithfulness. Mean scores were also close: Gemini support increased from 95.03% for zero-shot to 96.38% for structured reasoning, and MiniCheck support from 85.42% to 85.94%, but these small descriptive differences were not supported as prompt main effects by the inferential models.")
    add_body(doc, "This result is credible in light of the experimental control. All prompts requested the same task, audience, fidelity and concision requirements; the role prompt added a specialist persona, while the structured prompt added an internal checking procedure. The prompts therefore manipulated strategy without giving one condition a privileged content checklist. The absence of a broad structured-reasoning advantage suggests that techniques developed for multi-step reasoning tasks do not necessarily transfer to whole-document summarisation. Chain-of-thought prompting has produced strong gains on arithmetic and symbolic reasoning (Wei et al., 2022), but summarisation requires selection, compression and faithful paraphrase rather than recovery of a single reasoned answer.")
    add_body(doc, "Prompting did affect length and compression. Structured outputs averaged 644.03 words compared with 604.44 for zero-shot and 618.33 for role prompting, and the prompt effect on compression was statistically significant. The practical gain is therefore questionable: structured reasoning tended to consume more output space without delivering a reliable main effect on readability or faithfulness. For this task and these model versions, the simpler zero-shot prompt may be the preferable operational default because it performs comparably on the measured quality outcomes while producing shorter summaries. This conclusion is deliberately bounded to the exact prompts and API models used; stronger prompt engineering, examples or topic checklists would constitute different interventions and could produce different results.")

    add_heading(doc, "5.4 RQ3: Model-prompt interaction", 1)
    add_body(doc, "RQ3 asked whether the effect of prompting strategy varied across LLMs. The answer is qualified. A statistically significant interaction occurred for Gemini faithfulness, summary length and compression, but not for MiniCheck faithfulness or readability. Under Gemini, GPT exceeded Llama in the zero-shot and role conditions, and Mistral exceeded Llama in those same conditions; the corresponding structured-prompt contrasts were not significant. One possible interpretation is that structured reasoning narrowed Gemini-assessed differences between models. However, because MiniCheck did not reproduce the interaction, this should be described as evaluator-specific rather than a general improvement in factual support.")
    add_body(doc, "The interaction results reinforce the danger of selecting a prompt in isolation from the model. A prompt that changes verbosity for one model may have little effect for another, and a judge may reward particular formulations or evidence patterns. Nevertheless, the lack of interaction for readability means Mistral's readability advantage was comparatively stable across prompts, while the absence of a MiniCheck interaction indicates that Llama's advantage under that evaluator was not dependent on a particular prompt. RQ3 is therefore answered by partial, outcome-specific interaction rather than a single universal crossover effect.")

    add_heading(doc, "5.5 Evaluator sensitivity and human validation", 1)
    add_body(doc, "The multi-evaluator design is one of the study's main methodological strengths. Conventional overlap metrics do not directly assess whether generated content is supported by the source, and factual inconsistency remains a distinct risk in abstractive summarisation (Wang, Cho and Lewis, 2020; Kryscinski et al., 2020). Claim-level evaluation also made errors inspectable by linking judgements to specific summary claims and source evidence rather than assigning only one opaque score per summary.")
    add_body(doc, "At the same time, the results demonstrate that an automated evaluator is not a neutral measuring instrument. Human-Gemini strict agreement was high, but Cohen's kappa was only 0.231 because supported claims dominated the sample. This is a known prevalence problem: high raw agreement can coexist with a low chance-corrected coefficient when one category is overwhelmingly common. The dissertation should therefore report both percentage agreement and kappa, rather than using either in isolation. The hybrid human sample, which combined broad representation with additional difficult claims, was more informative than a purely random sample would have been, but it also means the human percentages are validation statistics and should not be presented as population estimates for all 10,485 claims.")
    add_body(doc, "The most defensible use of the automated results is triangulation. Gemini provides the judge most closely aligned with the human sample; MiniCheck supplies an independent architecture and a stricter counter-view; and human annotation provides direct adjudicative evidence on a manageable subset. Where Gemini and MiniCheck agree, confidence is increased. Where they disagree, the study should preserve the disagreement and, if the claim is important to a practical decision, prioritise human review.")

    add_heading(doc, "5.6 Practical implications", 1)
    add_body(doc, "The findings support a configurable rather than one-size-fits-all summarisation pipeline. For public deployment, model selection should follow the communication objective:")
    add_bullet(doc, "Prefer Mistral when readability is the primary objective and a longer first draft is acceptable.")
    add_bullet(doc, "Prefer Llama when strong compression and favourable MiniCheck grounding are prioritised, while reviewing linguistic complexity.")
    add_bullet(doc, "Prefer GPT when the Gemini-aligned faithfulness signal is prioritised, while accepting lower readability than Mistral.")
    add_body(doc, "These are not deployment endorsements without qualification. Privacy summaries can influence users' understanding of rights, retention, sharing and automated processing. A safe operational design would retain links to the full policy, expose supporting evidence for important claims, flag evaluator disagreements, and use human review for high-risk or public-facing outputs. Regulatory guidance also recommends layered presentation and user testing, which cannot be replaced by a Flesch score alone (Information Commissioner's Office, n.d.-b). The Streamlit evidence interface developed around the study is therefore methodologically useful: it turns support labels into reviewable evidence rather than treating the aggregate percentage as the final product.")

    add_heading(doc, "5.7 Limitations", 1)
    add_body(doc, "Several limitations bound the conclusions. First, the study compares three specific API model versions, not the model families in perpetuity. Provider updates, hidden system behaviour and decoding differences can alter results. Second, one generation was retained per main experimental condition. The pilot reproducibility check supported this economical choice, but a repeated-generation design would estimate within-condition variability more directly. Third, the 30-policy corpus was purposively stratified across six sectors; it improves document diversity but is not a random sample of all privacy policies, and sector was not powered as an inferential factor.")
    add_body(doc, "Fourth, Flesch Reading Ease captures sentence and word characteristics, not actual comprehension, navigation, trust or legal accuracy. The finding that Mistral was more readable should therefore be tested with non-specialist users before making a usability claim. Fifth, claim decomposition and automated support judgements introduce measurement choices. The high class imbalance limits kappa, the human sample does not cover every claim, and the Gemini-MiniCheck divergence shows that conclusions can depend on evaluator design. Sixth, the full main-study information-coverage annotation was not completed. The study consequently cannot determine whether the most faithful or most concise system also retained the greatest proportion of important source content.")
    add_body(doc, "Finally, no single composite quality score was calculated. This was appropriate because weighting faithfulness, readability and conciseness would embed an arbitrary value judgement. The trade-off profile is more transparent, although it prevents a simple league table.")

    add_heading(doc, "5.8 Recommendations for future work", 1)
    add_body(doc, "Future research should repeat each model-prompt-policy condition to estimate stochastic variability and should replicate the experiment as models are updated. A complete source-unit coverage study would directly address the remaining part of RQ1 and clarify whether compression removes legally material qualifications. Human evaluation should expand beyond researchers to include non-specialist readers, measuring comprehension, perceived clarity, decision confidence and time-on-task. It would also be valuable to test evidence-linked or retrieval-assisted summarisation, automatic post-generation verification, and targeted revision of claims on which evaluators disagree.")
    add_body(doc, "The present interaction findings also motivate adaptive prompting. Rather than applying one prompt to every model, future systems could select constraints based on a model's observed failure mode—for example, enforcing a length ceiling for Mistral, requesting plainer sentence structure from Llama, or adding a coverage checklist after independently defining the required privacy concepts. Such experiments should keep the evaluation taxonomy separate from prompt design when the aim is an unbiased comparison, or clearly redefine the study as an optimisation task when using that taxonomy to improve outputs.")

    add_heading(doc, "5.9 Overall answer to the research question", 1)
    add_body(doc, "The study shows that selected LLM choice materially affects plain-English privacy-policy summarisation, but the direction of advantage depends on the quality dimension and faithfulness evaluator. Prompt strategy, when task requirements are tightly controlled, has a weaker and less reliable independent effect. Structured reasoning slightly increased output length and did not deliver a general readability or faithfulness benefit. Model-prompt interactions exist for some outcomes, particularly Gemini-assessed support and compression, but are not universal. The most defensible conclusion is therefore a trade-off: Mistral offers greater readability at lower compression, Llama offers greater conciseness and stronger MiniCheck scores, and GPT offers stronger Gemini-aligned support. Human validation makes Gemini the more credible primary faithfulness indicator in this dataset, while the disagreement with MiniCheck remains an important uncertainty rather than an error to be averaged away.")

    add_heading(doc, "References", 1)
    refs = [
        "Amos, R., Acar, G., Lucherini, E., Kshirsagar, M., Narayanan, A. and Mayer, J. (2022) 'Privacy policies across the ages: Content and readability of privacy policies 1996-2021', Proceedings on Privacy Enhancing Technologies, 2022(1). Available at: https://arxiv.org/abs/2201.08739",
        "Bhandari, M., Gour, P.N., Ashfaq, A. and Liu, P. (2020) 'Metrics also disagree in the low scoring range: Revisiting summarization evaluation metrics', Proceedings of COLING 2020. Available at: https://aclanthology.org/2020.coling-main.501/",
        "Information Commissioner's Office (n.d.-a) The right to be informed. Available at: https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/individual-rights/the-right-to-be-informed/ (Accessed: 20 August 2026).",
        "Information Commissioner's Office (n.d.-b) How should we draft our privacy information? Available at: https://ico.org.uk/for-organisations/uk-gdpr-guidance-and-resources/individual-rights/the-right-to-be-informed/how-should-we-draft-our-privacy-information/ (Accessed: 20 August 2026).",
        "Kryscinski, W., McCann, B., Xiong, C. and Socher, R. (2020) 'Evaluating the factual consistency of abstractive text summarization', Proceedings of EMNLP 2020, pp. 9332-9346. doi:10.18653/v1/2020.emnlp-main.750.",
        "Mujahid, Z.M., Wright, D. and Augenstein, I. (2026) 'Stress testing factual consistency metrics for long-document summarization', Proceedings of ACL 2026, pp. 31914-31933. doi:10.18653/v1/2026.acl-long.1472.",
        "Tang, L., Laban, P. and Durrett, G. (2024) 'MiniCheck: Efficient fact-checking of LLMs on grounding documents', Proceedings of EMNLP 2024, pp. 8818-8847. doi:10.18653/v1/2024.emnlp-main.499.",
        "Wang, A., Cho, K. and Lewis, M. (2020) 'Asking and answering questions to evaluate the factual consistency of summaries', Proceedings of ACL 2020, pp. 5008-5020. doi:10.18653/v1/2020.acl-main.450.",
        "Wei, J. et al. (2022) 'Chain-of-thought prompting elicits reasoning in large language models', Advances in Neural Information Processing Systems, 35. Available at: https://proceedings.neurips.cc/paper_files/paper/2022/hash/9d5609613524ecf4f15af0f7b31abca4-Abstract-Conference.html",
    ]
    for ref in refs:
        p = doc.add_paragraph(ref)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_after = Pt(6)
        for r in p.runs:
            font(r, size=10)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
