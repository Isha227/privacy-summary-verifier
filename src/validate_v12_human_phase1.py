from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import load_workbook

from .common import ROOT, read_csv


LABELS = {"Supported", "Partially supported", "Contradicted", "Unsupported"}
FAILURES = {"None", "Distortion", "Hallucination", "Hallucination and distortion"}
CONFIDENCE = {"High", "Medium", "Low"}


def _records(sheet) -> list[dict]:
    values = list(sheet.values)
    headers = [str(value or "").strip() for value in values[0]]
    return [dict(zip(headers, row)) for row in values[1:]]


def validate_workbook(path: Path, evaluator: str) -> tuple[list[dict], list[dict], list[str]]:
    expected = {row["statement_id"]: row for row in read_csv(
        ROOT / "data/v2_prompt_intervention_v1_2/minicheck_colab/v12_minicheck_statements.csv"
    )}
    workbook = load_workbook(path, data_only=False, read_only=True)
    errors: list[str] = []
    if "Faithfulness" not in workbook.sheetnames or "Source Units" not in workbook.sheetnames:
        return [], [], ["Required Faithfulness or Source Units sheet is missing"]
    faith = _records(workbook["Faithfulness"])
    if len(faith) != 1393:
        errors.append(f"Faithfulness row count is {len(faith)}, expected 1393")
    seen = set()
    for row_number, row in enumerate(faith, 2):
        statement_id = str(row.get("statement_id") or "").strip()
        if statement_id in seen: errors.append(f"Duplicate statement ID {statement_id}")
        seen.add(statement_id)
        exp = expected.get(statement_id)
        if not exp:
            errors.append(f"Unexpected statement ID at row {row_number}: {statement_id}"); continue
        for field in ("policy_id", "blind_id", "statement_text"):
            if str(row.get(field) or "") != str(exp[field]):
                errors.append(f"Frozen {field} changed for {statement_id}")
        label = str(row.get("label") or "").strip()
        failure = str(row.get("failure_type") or "").strip()
        if label not in LABELS: errors.append(f"Invalid/missing label for {statement_id}")
        if failure not in FAILURES: errors.append(f"Invalid/missing failure type for {statement_id}")
        if label == "Supported" and failure != "None": errors.append(f"Supported must use None: {statement_id}")
        if label in LABELS - {"Supported"} and failure == "None": errors.append(f"Non-supported cannot use None: {statement_id}")
        if label in {"Supported", "Partially supported", "Contradicted"} and not str(row.get("exact_source_evidence") or "").strip():
            errors.append(f"Missing source evidence for {statement_id}")
        if not str(row.get("reason") or "").strip(): errors.append(f"Missing reason for {statement_id}")
        if str(row.get("confidence") or "").strip() not in CONFIDENCE: errors.append(f"Invalid/missing confidence for {statement_id}")
        row["evaluator"] = evaluator
    if seen != set(expected): errors.append(f"Statement ID set differs: missing={len(set(expected)-seen)} extra={len(seen-set(expected))}")

    units = []
    for row in _records(workbook["Source Units"]):
        include = str(row.get("include") or "").strip()
        if include not in {"Yes", "No", ""}: errors.append(f"Invalid source-unit include value: {include}")
        if include != "Yes": continue
        unit_id = str(row.get("proposed_unit_id") or "").strip()
        for field in ("policy_id", "proposed_unit_id", "exact_source_quote", "important_information_in_plain_language", "why_important_to_user"):
            if not str(row.get(field) or "").strip(): errors.append(f"Included unit {unit_id or '[missing ID]'} lacks {field}")
        row["evaluator"] = evaluator
        units.append(row)
    return faith, units, errors


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--root", default="data/v2_prompt_intervention_v1_2/human_evaluation")
    args = parser.parse_args()
    base = ROOT / args.root
    report = {"status": "valid", "evaluators": {}}
    all_errors = []
    for evaluator in ("A1", "A2", "A3"):
        package = base / f"Researcher_{evaluator}_Package"
        candidates = sorted(package.glob(f"human_v12_{evaluator}_PHASE1_COMPLETED.xlsx"))
        if len(candidates) != 1:
            errors = [f"Expected one completed workbook; found {len(candidates)}"]
            faith, units = [], []
        else:
            faith, units, errors = validate_workbook(candidates[0], evaluator)
        report["evaluators"][evaluator] = {"faithfulness_rows": len(faith), "included_source_units": len(units), "errors": errors}
        all_errors.extend(f"{evaluator}: {error}" for error in errors)
    report["status"] = "valid" if not all_errors else "invalid"
    output = base / "phase1_validation_report.json"
    output.write_text(json.dumps(report, indent=2), encoding="utf-8")
    print(output)
    if all_errors:
        print(f"INVALID: {len(all_errors)} problems found")
        raise SystemExit(1)
    print("VALID: all three Phase 1 workbooks are complete and structurally frozen")


if __name__ == "__main__":
    main()
